from common import *
from analyze import SHORT,NAMES,DEFAULT,TARGETS
import sys
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter

COLORS={'gt10_hard':'#172334','gt10_soft':'#8055b3','e29_cost_margin':'#909aa6','e29_cost_ce':'#bd7180',
    'single':'#2975b6','multi':'#e48022','multi_matched':'#23966c','single_matched':'#269d9d',
    'e28_mlp':'#d95366','nsg':'#7d5b3e','fast30k':'#6e9c42','gt1_hard':'#b59a23','gt1_soft':'#48a5b0'}
STYLES={'gt10_hard':'-','gt10_soft':'--','e29_cost_margin':':','e29_cost_ce':'-.','single':'--','multi':'-','multi_matched':'-.'}
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.titlesize':14,'axes.labelsize':12,
    'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none','savefig.facecolor':'white'})

def save(fig,name,caption='',bottom=.17):
    fig.subplots_adjust(left=.075,right=.985,top=.88,bottom=bottom,wspace=.23,hspace=.45)
    if caption:fig.text(.5,.018,caption,ha='center',va='bottom',fontsize=9,color='#58616b')
    fig.savefig(O/'figures'/f'{name}.png',dpi=160);fig.savefig(O/'figures'/f'{name}.svg');plt.close(fig)

def legend(fig,axs,ncol=3,y=.07):
    handles={}
    for ax in np.ravel(axs):
        hs,ls=ax.get_legend_handles_labels()
        for h,l in zip(hs,ls):handles.setdefault(l,h)
    fig.legend(handles.values(),handles.keys(),loc='lower center',bbox_to_anchor=(.5,y),ncol=ncol,frameon=False,fontsize=10)

def qps_axis(ax,rows):
    vals=[v for r in rows for v in [r['qps_min'],r['qps_max']]];lo=min(vals)*.90;hi=max(vals)*1.10
    ax.set_yscale('log');ax.set_ylim(lo,hi)
    ticks=[x for x in [100,200,400,800,1000,2000,4000,8000,10000] if lo<=x<=hi]
    ax.set_yticks(ticks);ax.yaxis.set_major_formatter(FuncFormatter(lambda v,p:f'{v:,.0f}'));ax.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v,p:f'{v:g}%'))
    ax.set_xlabel('Recall@10');ax.set_ylabel('Queries / second (log scale)');ax.grid(alpha=.18)

def curve_figure(tables,name,all_methods=False):
    fig,axs=plt.subplots(1,2,figsize=(16,7.3 if all_methods else 6.6))
    for ax,ds in zip(axs,DS):
        rows=[r for r in tables['all_points'] if r['dataset']==SHORT[ds] and r['stage']=='validation']
        keys=list(dict.fromkeys(r['method'] for r in rows))
        if not all_methods:keys=[k for k in ['gt10_hard',DEFAULT[ds],'e29_cost_margin','single','multi','multi_matched','single_matched'] if k in keys]
        keys=list(dict.fromkeys(keys))
        for key in keys:
            rs=sorted([r for r in rows if r['method']==key],key=lambda r:r['L'])
            y=np.array([r['qps'] for r in rs]);low=np.array([r['qps_min'] for r in rs]);high=np.array([r['qps_max'] for r in rs])
            ax.errorbar([r['recall_pct'] for r in rs],y,yerr=[y-low,high-y],label=NAMES[key],color=COLORS[key],
                linestyle=STYLES.get(key,'--'),marker='o',markersize=4,linewidth=1.7,capsize=2,alpha=.92)
        qps_axis(ax,rows);ax.set_title(SHORT[ds])
    fig.suptitle('Recall–QPS: all compared methods' if all_methods else 'Recall–QPS: equivalent teachers and controls',fontsize=18,y=.97)
    legend(fig,axs,ncol=4 if all_methods else 3,y=.06)
    save(fig,name,'L = 64, 128, 256, 400, 800. Median of 5 single-thread batch timings; whiskers show the full timing range.',bottom=.26 if all_methods else .21)

def selection_figure(tables):
    fig,axs=plt.subplots(1,2,figsize=(16,7.6))
    selected=['gt10_hard','gt10_soft','gt1_soft','fast30k','e28_mlp','e29_cost_ce','e29_cost_margin','single','multi','multi_matched','single_matched']
    for ax,ds in zip(axs,DS):
        rows=[r for r in tables['target_ratios'] if r['dataset']==SHORT[ds] and r['covered']]
        for key in selected:
            rs=sorted([r for r in rows if r['method']==key and key!=DEFAULT[ds]],key=lambda r:r['recall_pct'])
            if not rs:continue
            y=np.array([r['change_pct'] for r in rs]);lo=np.array([(r['ratio_min']-1)*100 for r in rs]);hi=np.array([(r['ratio_max']-1)*100 for r in rs])
            # Individual repeat ratios need not enclose ratio of the two medians.
            ax.plot([r['recall_pct'] for r in rs],y,color=COLORS[key],marker='o',linestyle=STYLES.get(key,'--'),label=NAMES[key],linewidth=1.6)
            ax.vlines([r['recall_pct'] for r in rs],lo,hi,color=COLORS[key],alpha=.4,linewidth=1)
        ax.axhline(0,color='#172334',linewidth=.9);ax.axhline(3,color='#23966c',linestyle=':',linewidth=1);ax.axhline(-3,color='#d95366',linestyle=':',linewidth=1)
        ax.set_xticks([x*100 for x in TARGETS[ds]]);ax.set_xlabel('Reference Recall@10 (%)');ax.set_ylabel('Interpolated QPS change (%)')
        ax.set_title(SHORT[ds]+' — reference: '+NAMES[DEFAULT[ds]]);ax.grid(alpha=.15)
    fig.suptitle('Speed at the frozen recall targets',fontsize=18,y=.97);legend(fig,axs,ncol=4,y=.055)
    save(fig,'selection_targets','Recall–logQPS interpolation on measured Pareto segments. Vertical ranges use the 5 timing rounds; they are not confidence intervals.',bottom=.27)

def training_figure(tables):
    fig,axs=plt.subplots(2,2,figsize=(15,8.3))
    for col,ds in enumerate(DS):
        for key in ['single','multi']:
            rs=sorted([r for r in tables['training_alignment'] if r['dataset']==SHORT[ds] and r['method']==key and r['teacher']=='multi' and r['group']=='all'],key=lambda r:r['epoch'])
            for row,field in enumerate(['candidate_acceptance','full_exact_accepted_fraction']):
                ax=axs[row,col];ax.plot([r['epoch'] for r in rs],[r[field]*100 for r in rs],marker='o',label=NAMES[key],color=COLORS[key],linestyle=STYLES[key])
                ax.set_xlabel('Epoch');ax.set_ylabel('Accepted complete triples (%)');ax.grid(alpha=.18)
        axs[0,col].set_title(SHORT[ds]+' — best of 56 labeled triples')
        axs[1,col].set_title(SHORT[ds]+' — exact top3 over all hubs')
    fig.suptitle('Training alignment to the same equivalent teacher sets',fontsize=18,y=.98);legend(fig,axs,ncol=2,y=.052)
    save(fig,'training_acceptance','Both models are scored against the same G. Exact score acceptance is a training diagnostic, not NSG Recall or QPS.',bottom=.16)

def transfer_figure(tables):
    fig,axs=plt.subplots(1,2,figsize=(16,6.8));stages=['fit','cal','audit','validation'];keys=['e29_cost_margin','single','multi','multi_matched']
    for ax,ds in zip(axs,DS):
        for idx,key in enumerate(keys):
            rs=[next(r for r in tables['delta_gt10'] if r['dataset']==SHORT[ds] and r['stage']==stage and r['method']==key and r['L']==128 and not r['exact']) for stage in stages]
            x=np.arange(4)+(idx-1.5)*.12;y=np.array([r['delta_recall_pp'] for r in rs]);lo=np.array([r['ci_low_pp'] for r in rs]);hi=np.array([r['ci_high_pp'] for r in rs])
            ax.errorbar(x,y,yerr=[y-lo,hi-y],marker='o',capsize=3,color=COLORS[key],linestyle=STYLES[key],linewidth=1.4,label=NAMES[key])
        ax.axhline(0,color='#172334',linewidth=1);ax.set_xticks(range(4),['Fit 20k','Cal 1,024','Audit 1,024','Validation 2k'])
        ax.set_ylabel('Recall change vs GT10 hard (pp)');ax.set_title(SHORT[ds]);ax.grid(axis='y',alpha=.18)
    fig.suptitle('Actual HNSW → NSG: training-to-evaluation transfer at L128',fontsize=17,y=.97);legend(fig,axs,ncol=4,y=.07)
    save(fig,'fit_to_validation','Each split is compared with its own GT10 baseline. Whiskers: paired-query bootstrap 95% intervals, without multiplicity correction.',bottom=.20)

def drift_figure(tables):
    fig,axs=plt.subplots(2,2,figsize=(15,8));keys=['single','multi','multi_matched'];labels=['Single','Multi','Multi matched']
    for col,ds in enumerate(DS):
        g=next(r for r in tables['drift_match'] if r['dataset']==SHORT[ds]);vals=[g['single_raw_drift'],g['multi_raw_drift'],g['actual_drift']]
        axs[0,col].bar(range(3),np.array(vals)*1e4,color=[COLORS[k] for k in keys],width=.6)
        axs[0,col].axhline(g['target_drift']*1e4,color='#172334',linestyle=':',linewidth=1);axs[0,col].set_ylabel('Mean squared displacement × 10⁴')
        axs[0,col].set_title(SHORT[ds]+f" — multi scale t = {g['alpha']:.4f}")
        rs=[next(r for r in tables['delta_gt10'] if r['dataset']==SHORT[ds] and r['stage']=='fit' and r['method']==k and not r['exact']) for k in keys]
        y=np.array([r['delta_recall_pp'] for r in rs]);lo=np.array([r['ci_low_pp'] for r in rs]);hi=np.array([r['ci_high_pp'] for r in rs])
        axs[1,col].bar(range(3),y,color=[COLORS[k] for k in keys],width=.6)
        axs[1,col].errorbar(range(3),y,yerr=[y-lo,hi-y],fmt='none',ecolor='#172334',capsize=4)
        axs[1,col].axhline(0,color='#172334',linewidth=1);axs[1,col].set_ylabel('Actual fit Recall change (pp)');axs[1,col].set_title('HNSW → original NSG, L128')
        for ax in axs[:,col]:ax.set_xticks(range(3),labels);ax.grid(axis='y',alpha=.15)
    fig.suptitle('Control the amount of profile change',fontsize=18,y=.98)
    save(fig,'drift_control','The single model is its own matched control. Mean displacement is matched; per-hub changes and routing decisions need not match.',bottom=.13)

def equivalence_figure(tables):
    fig,axs=plt.subplots(1,2,figsize=(16,6.3))
    for ax,ds in zip(axs,DS):
        for group,label,color in [('headroom','Queries with recall headroom','#e48022'),('no_headroom','Queries without recall headroom','#65778b')]:
            rs=[r for r in tables['equivalence_histogram'] if r['dataset']==SHORT[ds] and r['group']==group]
            ax.step([r['equivalent_sets'] for r in rs],[r['fraction']*100 for r in rs],where='mid',label=label,color=color,linewidth=1.8)
        ax.set_title(SHORT[ds]);ax.set_xlabel('Number of tied optimal triples (out of 56)');ax.set_ylabel('Queries within group (%)')
        ax.set_xticks([1,4,8,16,24,32,40,48,56]);ax.set_xlim(.5,56.5);ax.grid(alpha=.15)
    fig.suptitle('How many complete triples are equally good?',fontsize=18,y=.96);legend(fig,axs,ncol=2,y=.06)
    save(fig,'label_equivalence','Equal highest GT10 hit count AND equal minimum graph NDC at L128. Equal graph work is not equal wall time.',bottom=.20)

def workbook(tables):
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        ws=wb.create_sheet(name);fields=list(dict.fromkeys(k for row in rows for k in row));ws.append(fields)
        for row in rows:ws.append([row.get(k) for k in fields])
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23384D')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            letter=col[0].column_letter;ws.column_dimensions[letter].width=min(48,max(14,len(str(col[0].value))+2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float):cell.number_format='0.000000'
    ws=wb.create_sheet('Sources');ws.append(['Item','Source / definition'])
    for item,value in [('Protocol',str(O/'protocol.md')),('Raw outputs',str(O/'results')),('Report',str(O/'report.md')),
        ('QPS','Median of 5 independent single-thread end-to-end batch timings; each raw time is in all_points'),
        ('Recall','Recall@10 percent, raw IP, original NSG'),('Interpolation','Measured Pareto segments in Recall-logQPS; no extrapolation'),
        ('Intervals','Recall: paired query bootstrap; QPS ranges: observed 5 rounds; neither is blind confirmation'),
        ('Teacher','All C8 triples tied in highest hit count and minimum graph NDC at L128'),
        ('Alias','single_matched is exactly single; no duplicate independent measurement')]:ws.append([item,value])
    ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=115
    wb.save(O/'chart_data.xlsx')

if __name__=='__main__':
    assert read(O/'state.json')['status']=='DIAGNOSTICS_COMPLETE_PLOTS_PENDING'
    tables=read(O/'results/analysis_tables.json')
    curve_figure(tables,'recall_qps_main');curve_figure(tables,'recall_qps_all',True);selection_figure(tables)
    training_figure(tables);transfer_figure(tables);drift_figure(tables);equivalence_figure(tables);workbook(tables)
    dump(O/'state.json',dict(status='PLOTS_COMPLETE_REPORT_REVIEW_PENDING',training_runs=4,real_search_calls=NORMAL_CAP,heldout_access=False))
    print('PLOTS_AND_WORKBOOK_COMPLETE',flush=True)
