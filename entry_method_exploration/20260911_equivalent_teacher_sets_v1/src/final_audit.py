from common import *
import math, re, sys, xml.etree.ElementTree as ET
import numpy as np
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import load_workbook

FIGURES=['recall_qps_main','recall_qps_all','selection_targets','training_acceptance','fit_to_validation','drift_control','label_equivalence']

def links(text):return re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',text)

def scalar_interpolation(curve,target,repeat=None):
    nondominated=[]
    for a in curve:
        if any(b['recall_at_10']>=a['recall_at_10'] and b['qps']>=a['qps'] and
            (b['recall_at_10']>a['recall_at_10'] or b['qps']>a['qps']) for b in curve):continue
        nondominated.append(a)
    pts=sorted(nondominated,key=lambda x:x['recall_at_10'])
    if not pts or not pts[0]['recall_at_10']<=target<=pts[-1]['recall_at_10']:return None
    value=lambda p:p['qps'] if repeat is None else p['qps_runs'][repeat]
    for p in pts:
        if p['recall_at_10']==target:return value(p)
    for a,b in zip(pts,pts[1:]):
        if a['recall_at_10']<target<b['recall_at_10']:
            f=(target-a['recall_at_10'])/(b['recall_at_10']-a['recall_at_10'])
            return math.exp((1-f)*math.log(value(a))+f*math.log(value(b)))
    raise AssertionError('missing interpolation segment')

def main():
    dest=O/'reviews/final_audit.json';dump(dest,dict(status='RUNNING',started=time.time()))
    try:
        protected=read(O/'reviews/protected_before.json')
        for item in protected:assert identity(item['path'])==item,item['path']
        frozen=0
        for name in ['reviews/protocol_lock.json','reviews/frozen_before_training.json','models/raw_frozen_before_matching.json',
            'models/frozen_before_development.json','reviews/evaluation_start_freeze.json','models/frozen_before_validation.json']:
            check_freeze(O/name);frozen+=len(read(O/name)['files'])
        assert sha(R/'AGENTS.md')==sha(O/'source_before/AGENTS.md')
        for name in ['teacher_verification','objective_tests','model_verification','output_verification','set_diagnostics','plot_visual_review','report_review','document_update']:
            assert read(O/'reviews'/f'{name}.json')['status']=='PASS',name
        for item in read(O/'reviews/document_update.json')['files']:assert identity(item['path'])==item,item['path']
        for item in read(O/'reviews/plot_visual_review.json')['files']:assert identity(item['path'])==item,item['path']
        for item in read(O/'reviews/report_review.json')['files']:assert identity(item['path'])==item,item['path']
        old=(E29/'src/evaluate.cpp').read_text();new=(O/'src/evaluate.cpp').read_text()
        assert new.replace('(reps==0 || reps==5)','(reps==0 || reps==3)')==old and new!=old
        build=read(O/'bin/evaluate.build.json');assert build['exit_code']==0 and identity(O/'bin/evaluate')==build['binary']
        for n in ['benchmark_with_nsg.cpp','legacy_reference.cpp']:assert sha(O/'src'/n)==sha(E29/'src'/n)
        runs=read(O/'models/training_runs.json')
        assert len(runs)==4 and {(r['dataset'],r['cell']) for r in runs}=={(ds,cell) for ds in DS for cell in CELLS}
        for r in runs:
            assert r['epochs']==80 and r['seed']==SEED and r['lambda_anchor']==100 and r['temperature']==.05 and r['margin']==.1
            assert not any(r[k] for k in ['cal_used','audit_used','validation_used'])
            assert identity(r['W']['path'])==r['W'] and identity(r['A']['path'])==r['A']
        logical=0
        for ds in DS:
            td=O/'inputs'/ds/'teachers';n=20000
            hits=np.fromfile(E28/'labels'/ds/'hits.u8',np.uint8).reshape(n,56)
            ndc=np.fromfile(E28/'labels'/ds/'graph_ndc.u32',np.uint32).reshape(n,56)
            g=np.fromfile(td/'multi.u8',np.uint8).reshape(n,56).astype(bool)
            top=hits==hits.max(1,keepdims=True);minimum=np.where(top,ndc,np.iinfo(np.uint32).max).min(1)
            assert np.array_equal(g,top&(ndc==minimum[:,None]))
            single=np.fromfile(td/'single.u8',np.uint8).reshape(n,56).astype(bool)
            assert (single.sum(1)==1).all() and (g|~single).all()
            assert read(td/'verification.json')['logical_labels_verified']==n*56;logical+=n*56
            match=read(O/'models'/ds/'drift_match.json')
            assert len(match['canonical_models'])==3 and match['aliases']=={'single_matched':'single'}
            assert match['relative_error']<1e-4 and not match['cal_used'] and not match['audit_used'] and not match['validation_used']
        jobs=read(O/'logs/evaluation_jobs.json');assert len(jobs)==8 and all(j['status']=='COMPLETE' and j['exit_code']==0 for j in jobs)
        calls=sum(j['reserved_calls'] for j in jobs);assert calls==NORMAL_CAP and calls<HARD_CAP
        assert sum(j['seconds'] for j in jobs)<5400
        planned=read(O/'reviews/planned_matrix.json');assert planned['calls']==calls
        actual_points=0;actual_outputs=0;formal=0
        for j in jobs:
            assert j['seconds']<2000 and j['rss_peak_kib']<40*(1<<20) and identity(j['config']['path'])==j['config']
            root=O/'results'/f"{j['stage']}_{j['dataset']}";r=read(root/'results.json');c=read(root/'config.json');th=read(root/'thread_contract.json')
            assert r['physical_graph_calls']==j['reserved_calls'] and th['actual_openmp_threads']==th['process_threads']==1 and th['cpu']==15
            assert r['GT_opened_after_all_raw_search'] and r['threads']==c['threads']==1
            assert read(root/'independent_verification.json')['status']=='PASS'
            actual_points+=len(r['methods']);actual_outputs+=c['num_queries']*len(r['methods'])
            if j['stage']=='validation':
                formal+=len(r['methods']);assert r['timing_repeats']==5
                for m in r['methods']:
                    assert len(m['qps_runs'])==len(m['end_to_end_seconds'])==5
                    assert np.allclose(m['qps_runs'],c['num_queries']/np.array(m['end_to_end_seconds']),rtol=1e-12)
        out=read(O/'reviews/output_verification.json')
        assert actual_points==out['total_points']==180 and actual_outputs==out['total_query_outputs']==491200
        assert formal==out['formal_points']==120 and out['historical_regression_outputs']==180000 and out['physical_calls']==calls
        assert out['fit_native_cache_parity_outputs']>=40000
        sd=read(O/'reviews/set_diagnostics.json');assert sd['NSG_calls']==sd['optimizer_steps']==0
        tables=read(O/'results/analysis_tables.json');sel=read(O/'results/method_selection.json');recomputed=[];interpolations=0
        for ds,short,default,targets in [(DS[0],'WebVid2.5M','gt10_soft',[75.,80.,85.]),(DS[1],'Yandex10M IP','gt10_hard',[90.,95.,98.])]:
            raw=read(O/'results'/f'validation_{ds}'/'results.json')['methods'];curves={}
            for m in raw:curves.setdefault(m['model_name'],[]).append(m)
            for row in [r for r in tables['all_points'] if r['dataset']==short and r['stage']=='validation']:
                dominates=lambda m:m['recall_at_10']*100>=row['recall_pct'] and m['qps']>=row['qps'] and (m['recall_at_10']*100>row['recall_pct'] or m['qps']>row['qps'])
                assert row['pareto_all_methods_and_L']==(not any(dominates(m) for m in raw))
                assert row['pareto_same_L']==(not any(dominates(m) for m in raw if m['L']==row['L']))
            intervals=[r for r in tables['frontier_intervals'] if r['dataset']==short]
            assert abs(intervals[0]['recall_low_pct']-targets[0])<1e-10 and abs(intervals[-1]['recall_high_pct']-targets[-1])<1e-10
            for i,row in enumerate(intervals):
                left,right=row['recall_low_pct']/100,row['recall_high_pct']/100
                assert right>left
                if i:assert abs(row['recall_low_pct']-intervals[i-1]['recall_high_pct'])<1e-9
                knots=sorted({left,right}|{m['recall_at_10'] for m in raw if left<m['recall_at_10']<right})
                probes=knots+[(a+b)/2 for a,b in zip(knots,knots[1:])]
                for target in probes:
                    chosen=scalar_interpolation(curves[row['method']],target)
                    others=[scalar_interpolation(curve,target) for curve in curves.values()]
                    assert chosen is not None and chosen+1e-7>=max(v for v in others if v is not None),(short,row,target,chosen,others)
            for row in [r for r in tables['target_ratios'] if r['dataset']==short]:
                assert row['reference']==default and row['recall_pct'] in targets
                a=scalar_interpolation(curves[row['method']],row['recall_pct']/100);b=scalar_interpolation(curves[default],row['recall_pct']/100)
                assert row['covered']==(a is not None and b is not None)
                if row['covered']:
                    assert abs(a/b-row['ratio'])<1e-10 and abs(a-row['interpolated_qps'])<1e-7
                    repeat=[]
                    for i in range(5):
                        ratio=scalar_interpolation(curves[row['method']],row['recall_pct']/100,i)/scalar_interpolation(curves[default],row['recall_pct']/100,i)
                        assert abs(ratio-row[f'ratio_run{i+1}'])<1e-10;repeat.append(ratio);interpolations+=1
                    assert row['passing_target']==(a/b>=1.03 and all(x>1 for x in repeat))
                    assert row['regression_target']==(a/b<=.97)
            for rank in [r for r in tables['method_ranking'] if r['dataset']==short]:
                rows=[r for r in tables['target_ratios'] if r['dataset']==short and r['method']==rank['method'] and r['covered']]
                assert len(rows)==rank['covered_targets']
                if rows:
                    geo=math.prod(r['ratio'] for r in rows)**(1/len(rows));assert abs(geo-rank['geomean_ratio'])<1e-12
                eligible=len(rows)==3 and sum(r['passing_target'] for r in rows)>=2 and not any(r['regression_target'] for r in rows)
                assert eligible==rank['eligible_replacement']
            ranks=[r for r in tables['method_ranking'] if r['dataset']==short and r['all_targets_covered']]
            order=lambda r:(-r['geomean_ratio'],-r['worst_change_pct'],r['method'])
            eligible=sorted([r for r in ranks if r['eligible_replacement'] and r['method']!=default],key=order)
            decision=next(x for x in sel['decisions'] if x['dataset']==short)
            assert decision['descriptive_fastest']==sorted(ranks,key=order)[0]['method']
            assert decision['recommended']==(eligible[0]['method'] if eligible else default)
            recomputed.append(decision)
        for name in FIGURES:
            png=O/'figures'/f'{name}.png';svg=O/'figures'/f'{name}.svg'
            assert png.read_bytes().startswith(b'\x89PNG\r\n\x1a\n') and png.stat().st_size>10000
            assert ET.parse(svg).getroot().tag.endswith('svg')
        wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);sheet_counts={}
        for name,rows in tables.items():
            values=wb[name].values;headers=list(next(values));stored=list(values);assert len(stored)==len(rows),(name,len(stored),len(rows))
            for expected,actual in zip(rows,stored):
                for k,v in zip(headers,actual):
                    target=expected.get(k)
                    if isinstance(target,float):assert v is not None and abs(v-target)<1e-8,(name,k,target,v)
                    else:assert v==target,(name,k,target,v)
            sheet_counts[name]=len(rows)
        assert 'Sources' in wb.sheetnames;wb.close()
        documents=[p for p in O.glob('*.md')];roots=[R/n for n in ['summary.md','spec.md','s-target.md','s-experiment.md']]
        roots += [R/'experiment_history/19_equivalent_teacher_sets.md',R/'experiment_history/README.md'];link_checks=0
        for p in documents+roots:
            before=O/'source_before'/('history_README.md' if p==R/'experiment_history/README.md' else p.name)
            oldlinks=set(links(before.read_text())) if p in roots and before.exists() else set()
            for link in links(p.read_text()):
                if link in oldlinks or link.startswith(('http:','https:','#','mailto:')):continue
                target=link.split('#',1)[0].strip('<>');path=Path(target) if target.startswith('/') else p.parent/target
                assert path.exists(),(p,link);link_checks+=1
        check=subprocess.run(['git','diff','--check','--','summary.md','spec.md','s-target.md','s-experiment.md','experiment_history'],cwd=R,capture_output=True,text=True)
        assert check.returncode==0,check.stdout+check.stderr
        for source in read(O/'sources/inventory.json'):
            assert sha(O/'sources'/source['file'])==source['sha256']
        size=sum(p.stat().st_size for p in O.rglob('*') if p.is_file());free=shutil.disk_usage(O).free
        assert size<1<<30 and free>10*(1<<30)
        conclusion=read(O/'results/conclusion.json');assert conclusion['recommendations']==sel['decisions']
        audit=dict(status='PASS',protected_identity_checks=len(protected),frozen_identity_checks=frozen,training_runs=4,unique_new_W=6,
            reused_logical_labels=logical,new_label_calls=0,real_search_calls=calls,repair_calls=0,synthetic_calls=0,
            formal_points=formal,formal_repeats=5,all_points=actual_points,all_query_outputs=actual_outputs,historical_regression_outputs=180000,
            fit_native_cache_parity_outputs=out['fit_native_cache_parity_outputs'],independent_selection_interpolations=interpolations,
            recommendations=recomputed,figures=len(FIGURES),workbook_rows=sheet_counts,local_link_checks=link_checks,
            artifact_bytes=size,free_bytes=free,heldout_opened=False,conclusion=conclusion['status'],finished=time.time())
        dump(dest,audit);dump(O/'state.json',dict(status='COMPLETE',conclusion=conclusion['status'],training_runs=4,real_search_calls=calls,heldout_access=False,final_audit=str(dest)))
        print(json.dumps(audit),flush=True)
    except BaseException as error:
        dump(dest,dict(status='FAILED',error=repr(error),timestamp=time.time()));raise

if __name__=='__main__':main()
