from common import *
import ast,csv,re,sys
sys.path.insert(0,str(R/'entry_method_exploration/20260911_gt10_shared_transform_v1/src/vendor'))
from openpyxl import load_workbook
assert not (O/'reviews/final_audit.json').exists()
state=read(O/'state.json');assert state['status']=='DEVELOPMENT_AND_WALL_COMPLETE_ANALYSIS_PENDING'
manifest_names=['upstream','E35_preserved','before_labels','before_training','training_controller','models_after_training','deployment_implementation','development_sources','before_audit']
manifest_names += [f'inputs_{phase}_{ds}' for ds in DS for phase in ['fit_check','cal','audit']]+[f'before_wall_{ds}' for ds in DS]
for name in manifest_names:check(O/'reviews'/f'{name}.json')
prior={x['path']:x for x in read(E35/'reviews/upstream.json')['files']};matched=0
for x in read(O/'reviews/upstream.json')['files']:
 if x['path'] in prior:assert x['sha256']==prior[x['path']]['sha256'];matched+=1
oldstat=0
for man,prefix in [(E34/'reviews/frozen_completion.json',E34),(E34/'reviews/E33_preserved.json',E33)]:
 for x in read(man)['files']:
  p=Path(x['path'])
  if p.is_relative_to(prefix):
   st=p.stat();assert st.st_size==x['bytes'] and st.st_mtime_ns==x['mtime_ns'],p;oldstat+=1
old_docs={x['path']:x for x in read(E35/'reviews/frozen_completion.json')['files']}
for name in ['AGENTS.md','summary.md','spec.md','s-target.md','s-experiment.md']:
 assert sha(O/'source_before'/name)==old_docs[str(R/name)]['sha256']
assert sha(R/'AGENTS.md')==sha(O/'source_before/AGENTS.md')
for p in (O/'src').glob('*.py'):ast.parse(p.read_text())
jobs=read(O/'logs/jobs.json');assert len(jobs)==10
assert sum(x['actual'] for x in jobs)==280352 and sum(x['HNSW_calls'] for x in jobs)==62464 and sum(x['reserved'] for x in jobs)<=300000
for j in jobs:
 assert j['status']=='COMPLETE' and j['actual']<=j['reserved'] and j['peak_threads']==1 and j['peak_RSS_KiB']<=40*(1<<20) and j['seconds']<1200
 c=read(O/'results'/j['name']/'complete.json');assert c['status']=='COMPLETE' and c['NSG_calls']==j['actual'] and c['HNSW_calls']==j['HNSW_calls'] and c['threads']==1 and not c['GT_read']
tj=read(O/'reviews/training_job.json');tv=read(O/'reviews/training_verification.json');pf=read(O/'reviews/preflight.json');dv=read(O/'reviews/deployment_verification.json')
assert tj['exit_code']==0 and tj['NSG_calls']==0 and tj['peak_threads']==1 and tj['seconds']<7200
assert tv['status']==pf['status']==dv['status']=='PASS' and tv['optimizer_updates']==24000
assert sum(x['checkpoints'] for x in tv['runs'])==248 and sum(x['trajectory_episodes_verified'] for x in tv['runs'])==768000
assert max(x['native_numpy_score_max_error'] for x in dv['checks'])<1e-12 and len(dv['checks'])==38
for x in tv['baseline_sources']:assert sha(x['path'])==x['sha256']
assert read(O/'reviews/before_labels.json')['time']<min(x['started_at'] for x in jobs)
assert read(O/'reviews/before_training.json')['time']<read(O/'reviews/training_controller.json')['time']<read(O/'reviews/models_after_training.json')['time']
assert read(O/'reviews/models_after_training.json')['time']<min(x['started_at'] for x in jobs if x['name'].startswith('cal_'))
assert max(x['closed_at'] for x in jobs if x['name'].startswith('cal_'))<read(O/'reviews/before_audit.json')['time']<min(x['started_at'] for x in jobs if x['name'].startswith('audit_'))
points=read(O/'results/points.json');paired=read(O/'results/paired.json');query=read(O/'results/query_results.json');sel=read(O/'results/cal_selection.json');gate=read(O/'results/development_gate.json')
assert len(points)==38 and len(paired)==28 and len(query)==12288
assert gate['status']=='STOP_DUAL_DEVELOPMENT_GATE' and [x['development_pass'] for x in gate['datasets']]==[False,True]
assert not gate['full_frontier_run'] and not gate['default_replaced']
labels_verified=outputs_verified=0;raw={};querymap={(x['dataset'],x['phase'],x['method'],x['row_id']):x for x in query}
def intersect(ids,gt):
 result=np.empty(ids.shape[:-1],np.uint8)
 for i in range(len(gt)):
  truth=set(gt[i,:10].tolist())
  for idx in np.ndindex(ids.shape[1:-1]):result[(i,)+idx]=len(set(ids[(i,)+idx].tolist())&truth)
 return result
def mix(x):
 mask=(1<<64)-1;x=(x+0x9e3779b97f4a7c15)&mask;x=((x^(x>>30))*0xbf58476d1ce4e5b9)&mask;x=((x^(x>>27))*0x94d049bb133111eb)&mask;return x^(x>>31)
for ds in DS:
 p=O/'inputs'/ds;c=read(p/'config.json');hubs=np.fromfile(c['hub_ids'],np.int32);gt=np.fromfile(p/'gt.i64',np.int64).reshape(2048,100);l=O/'results'/('labels_'+ds)
 ids=np.fromfile(l/'ids.i32',np.int32).reshape(2048,56,10);pool=np.fromfile(l/'pool.i32',np.int32).reshape(2048,8);entry=np.fromfile(l/'entries.i32',np.int32).reshape(2048,56,3)
 assert (ids>=0).all() and (np.diff(np.sort(ids,axis=-1),axis=-1)>0).all() and (pool>=0).all() and (pool<len(hubs)).all()
 assert np.array_equal(entry,np.sort(hubs[pool[:,TRIP]],axis=-1))
 hits=intersect(ids,gt);z=np.load(l/'scored.npz');assert np.array_equal(hits,z['hits']);route=np.fromfile(l/'route_ndc.u32',np.uint32);ndc=np.fromfile(l/'graph_ndc.u32',np.uint32).reshape(2048,56)
 utility=hits.astype(float)-hits[:,0,None]-.1*((route[:,None]+ndc.astype(float))/(route+ndc[:,0])[:,None]-1);assert np.array_equal(utility,z['utility']);labels_verified+=hits.size
 sets=[np.fromfile(p/'row_ids.i64',np.int64)[:1792],np.fromfile(p/'row_ids.i64',np.int64)[1792:],np.fromfile(p/'cal/row_ids.i64',np.int64),np.fromfile(p/'audit/row_ids.i64',np.int64)]
 assert [len(x) for x in sets]==[1792,256,512,512]
 for a,b in itertools.combinations(sets,2):assert not np.intersect1d(a,b).size
 for phase in ['fit_check','cal','audit','wall']:
  ip=p/('audit' if phase=='wall' else phase);n=64 if phase=='fit_check' else 512;truth=np.fromfile(ip/'gt.i64',np.int64).reshape(n,100);rows=np.fromfile(ip/'row_ids.i64',np.int64)
  plist=[x for x in points if x['dataset']==ds and x['phase']==phase]
  basendc=(ndc[:n,0]+route[:n]) if phase=='fit_check' else None;basehits=hits[:n,0] if phase=='fit_check' else None
  if phase!='fit_check':
   bdir=O/'results'/f'{phase}_{ds}/base';basehits=intersect(np.fromfile(bdir/'ids.i32',np.int32).reshape(n,10),truth);basendc=np.fromfile(bdir/'graph_ndc.u32',np.uint32).astype(float)+np.fromfile(bdir/'route_ndc.u32',np.uint32)
  for pt in plist:
   method=pt['method'];f=O/'results'/f'{phase}_{ds}'/method
   ids2=np.fromfile(f/'ids.i32',np.int32).reshape(n,10);dist=np.fromfile(f/'distances.f32',np.float32).reshape(n,10);s=np.fromfile(f/'states.i32',np.int32).reshape(n,4);act=np.fromfile(f/'actions.i32',np.int32).reshape(n,3);pool2=np.fromfile(f/'pool.i32',np.int32).reshape(n,8);ent=np.fromfile(f/'entries.i32',np.int32).reshape(n,3)
   rn=np.fromfile(f/'route_ndc.u32',np.uint32);gn=np.fromfile(f/'graph_ndc.u32',np.uint32);scores=np.fromfile(f/'scores.f64',np.float64).reshape(n,56);hh=intersect(ids2,truth);zz=np.load(f/'scored.npz');uu=hh.astype(float)-basehits-.1*((rn.astype(float)+gn)/basendc-1)
   assert np.isfinite(dist).all() and (ids2>=0).all() and (np.diff(np.sort(ids2,axis=1),axis=1)>0).all()
   assert np.array_equal(hh,zz['hits']) and np.allclose(uu,zz['utility'],rtol=0,atol=1e-14)
   assert abs(pt['recall_pct']-hh.mean()*10)<1e-12 and abs(pt['failure_pct']-(hh==0).mean()*100)<1e-12
   assert abs(pt['delta_recall_pp']-(hh.astype(float)-basehits).mean()*10)<1e-12
   assert pt['helped']==int((hh>basehits).sum()) and pt['harmed']==int((hh<basehits).sum())
   assert pt['rescued']==int(((basehits==0)&(hh>0)).sum()) and pt['new_failures']==int(((basehits>0)&(hh==0)).sum())
   assert pt['route_ndc']==float(rn.mean()) and pt['graph_ndc']==float(gn.mean()) and pt['mean_swaps']==float((act>0).sum(axis=1).mean())
   assert np.all(s[:,0]==0)
   if method!='base':
    assert (pool2>=0).all() and (pool2<len(hubs)).all() and (np.diff(np.sort(pool2,axis=1),axis=1)>0).all()
    assert np.array_equal(ent,np.sort(hubs[pool2[np.arange(n)[:,None],TRIP[s[:,-1]]]],axis=1))
    for i in range(n):
     alive=True;current=0
     for step in range(3):
      if not alive:assert act[i,step]==-1 and s[i,step+1]==current;continue
      expected=0
      if method=='random':expected=mix(int(rows[i])^2026091719^((step*0xd1b54a32d192ed03)&((1<<64)-1)))%16
      else:
       best=scores[i,current]
       for j in range(1,16):
        if scores[i,NEXT[current,j]]>best+1e-9:best=scores[i,NEXT[current,j]];expected=j
      assert act[i,step]==expected;current=int(NEXT[current,expected]);assert s[i,step+1]==current;alive=expected!=0
   else:assert np.all(s==0) and np.all(act==-1)
   raw[(ds,phase,method)]=dict(hits=hh,ndc=rn.astype(float)+gn,utility=uu)
   if phase in ['cal','audit']:
    for i,row in enumerate(rows):
     q=querymap[(ds,phase,method,int(row))];assert q['hits']==int(hh[i]) and q['failed']==bool(hh[i]==0) and q['graph_ndc']==int(gn[i]) and q['route_ndc']==int(rn[i]) and q['final_state']==int(s[i,-1]) and q['swaps']==int((act[i]>0).sum())
   if phase=='fit_check':
    assert np.array_equal(ids2,ids[np.arange(n),s[:,-1]]) and np.array_equal(gn,ndc[np.arange(n),s[:,-1]])
   if phase=='wall':
    for fn in ['ids.i32','distances.f32','entries.i32','states.i32','actions.i32','graph_ndc.u32','route_ndc.u32']:assert sha(f/fn)==sha(O/'results'/f'audit_{ds}'/method/fn)
    ts=pt['timings'];qps=np.array([x['QPS'] for x in ts]);assert len(ts)==5 and [x['repeat'] for x in ts]==list(range(5))
    assert np.allclose(qps,[n/x['wall_seconds'] for x in ts],rtol=1e-14) and pt['QPS_median']==float(np.median(qps)) and pt['QPS_min']==float(qps.min()) and pt['QPS_max']==float(qps.max())
   outputs_verified+=n
 for arm in ARMS:
  candidates=[x for x in points if x['dataset']==ds and x['phase']=='cal' and x['method'].startswith(arm+'_')];winner=sorted(candidates,key=lambda x:(-x['utility'],x['total_ndc'],x['method']))[0]
  assert winner['method']==sel[ds][arm]['method']==arm+'_20260918'
assert labels_verified==229376 and outputs_verified==15872
for b in paired:
 a=raw[(b['dataset'],b['phase'],b['method'])];z=raw[(b['dataset'],b['phase'],b['reference'])];delta=(a['hits'].astype(float)-z['hits'])*10
 assert b['delta_recall_pp']==float(delta.mean()) and b['helped']==int((delta>0).sum()) and b['harmed']==int((delta<0).sum())
 assert abs(b['utility_delta']-(a['utility']-z['utility']).mean())<1e-14
 rng=np.random.default_rng(20260920);ci=np.quantile(delta[rng.integers(0,512,(5000,512))].mean(axis=1),[.025,.975]);assert np.allclose(ci,[b['recall_CI_low_pp'],b['recall_CI_high_pp']],rtol=0,atol=1e-13)
csvfiles=['models','training_checkpoints','training_curves','points','paired','timings','query_results','swap_outcomes','swap_counts','action_counts'];csvrows=0;excelcells=0
wb=load_workbook(O/'results/query_swap_policy.xlsx',read_only=True,data_only=False)
for name in csvfiles:
 expected=read(O/'results'/f'{name}.json')
 with (O/'results'/f'{name}.csv').open() as f:actual=list(csv.DictReader(f))
 assert len(actual)==len(expected)
 for a,b in zip(actual,expected):
  for key,v in a.items():assert v==('' if b.get(key) is None else str(b[key])),(name,key,v,b.get(key))
 ws=wb[name[:31]];sheetrows=ws.iter_rows(values_only=True);header=next(sheetrows);assert list(header)==list(actual[0])
 for a,b in zip(actual,sheetrows):
  assert len(b)==len(header)
  for key,v in zip(header,b):
   source=a[key]
   if source=='':assert v is None
   elif source in ['True','False']:assert v==(source=='True')
   elif isinstance(v,(float,int)):assert abs(float(source)-v)<=1e-13*max(1,abs(float(source)))
   else:assert v==source
   excelcells+=1
 assert ws.max_row==len(actual)+1;csvrows+=len(actual)
for x in read(O/'reviews/workbook_sources.json'):assert sha(O/x['path'])==x['sha256']
wb.close()
assert read(O/'reviews/visual_review.json')['status']=='PASS'
for name in ['training_progress','development_quality_and_cost','swap_behavior']:
 for ext in ['png','svg','pdf']:assert (O/'figures'/f'{name}.{ext}').stat().st_size>1000
docs=[R/f for f in ['AGENTS.md','summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md','experiment_history/25_query_swap_policy.md']]+[O/f for f in ['README.md','report.md','protocol.md','method_basis.md']]
links=0
for p in docs:
 for link in re.findall(r'\]\(([^)]+)\)',p.read_text()):
  if '://' in link or link.startswith('#'):continue
  target=(p.parent/link.split('#')[0]).resolve();assert target.exists() or target==O/'reviews/final_audit.json',(p,link);links+=1
subprocess.run(['git','diff','--check','--','summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md'],cwd=R,check=True)
resources();artifact=sum(p.stat().st_size for p in O.rglob('*') if p.is_file());free=shutil.disk_usage(O).free
freeze(O/'reviews/frozen_completion.json',[p for p in O.rglob('*') if p.is_file() and p!=O/'state.json']+docs)
final=dict(status='PASS',new_NSG_calls=280352,HNSW_calls=62464,training_runs=8,optimizer_updates=24000,cached_RL_episodes_verified=768000,checkpoints_verified=248,
 labels_independently_set_scored=labels_verified,deployment_outputs_independently_set_scored=outputs_verified,native_deployment_points=38,wall_points=6,wall_repeats=30,all_actions_and_real_entries_verified=True,
 fit_cache_parity_queries=512,prior_GT10_baseline_parity_queries=4096,training_and_deployment_verification='PASS',all_prequery_pretraining_cal_selection_hashes_checked=True,
 prior_upstream_matches=matched,old_artifacts_stat_checked=oldstat,old_E35_entire_artifact_SHA_checked=True,validation_heldout_read_or_hash=False,CSV_JSON_rows=csvrows,Excel_cells_checked=excelcells,
 visual_review=True,document_links_checked=links,artifact_bytes_before_final_freeze=artifact,free_disk_bytes=free,default_replaced=False,full_frontier=False,closed_at=time.time())
dump(O/'reviews/final_audit.json',final)
state.update(status='COMPLETE_QUERY_SWAP_POLICY',conclusion=gate['conclusion'],development_gate=gate['status'],final_audit='PASS',closed_at=final['closed_at']);dump(O/'state.json',state)
print(json.dumps(final,indent=2),flush=True)
