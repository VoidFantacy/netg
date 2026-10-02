from common import *
import csv,sys
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator
sys.path.insert(0,str(R/'entry_method_exploration/20260911_gt10_shared_transform_v1/src/vendor'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
assert read(O/'state.json')['status']=='DEVELOPMENT_AND_WALL_COMPLETE_ANALYSIS_PENDING'
P=read(O/'results/points.json');C=read(O/'results/training_checkpoints.json');B=read(O/'results/paired.json');Q=read(O/'results/query_results.json')
plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none','savefig.dpi':160})
names={DS[0]:'WebVid',DS[1]:'Yandex-IP'}
labels={'base':'Original top3','random':'Random','rl_20260917':'RL seed17','rl_20260918':'RL seed18','exact_20260917':'Exact seed17','exact_20260918':'Exact seed18'}
colors={'base':'#34495e','random':'#9099a2','rl_20260917':'#71b6c7','rl_20260918':'#087f9c','exact_20260917':'#e9ae75','exact_20260918':'#ba611b'}
figdir=O/'figures';figdir.mkdir(exist_ok=True)
def save(fig,name,foot):
 fig.text(.04,.016,foot,fontsize=9,color='#45515c',va='bottom');fig.subplots_adjust(bottom=.12,top=.88,hspace=.38,wspace=.26)
 for ext in ['png','svg','pdf']:fig.savefig(figdir/f'{name}.{ext}',facecolor='white')
 plt.close(fig)
fig,axes=plt.subplots(2,2,figsize=(12,8.5))
for col,ds in enumerate(DS):
 for row,role in enumerate(['train','monitor']):
  ax=axes[row,col]
  for arm in ARMS:
   for seed in SEEDS:
    method=f'{arm}_{seed}';a=[x for x in C if x['dataset']==ds and x['arm']==arm and x['seed']==seed and x['role']==role]
    ax.plot([x['step'] for x in a],[x['delta_recall_pp'] for x in a],label=labels[method],c=colors[method],lw=1.7,ls='-' if seed==20260918 else '--')
    for x in a:
     if x['selected']:ax.scatter(x['step'],x['delta_recall_pp'],marker='*',s=130,color=colors[method],edgecolor='white',zorder=5)
  ax.axhline(0,c='#888888',lw=.8);ax.grid(alpha=.15);ax.set_title(f'{names[ds]}: {role} (n={1792 if role=="train" else 256})');ax.set_ylabel('Greedy Recall change (pp)');ax.set_xlabel('Optimizer update');ax.xaxis.set_major_locator(MaxNLocator(4))
fig.legend(*axes[0,0].get_legend_handles_labels(),loc='upper center',bbox_to_anchor=(.5,.948),ncol=4,fontsize=9)
fig.suptitle('Query-conditioned swap policy: all eight training runs',fontsize=16)
save(fig,'training_progress','Stars: checkpoint chosen by monitor utility with the fixed NDC gate. Final update = 3000.\nTrain/monitor are historical fit subsets; development results are measured separately.')
fig,axes=plt.subplots(2,2,figsize=(12,8.7))
methods=['random','rl_20260917','rl_20260918','exact_20260917','exact_20260918']
for col,ds in enumerate(DS):
 ax=axes[0,col]
 for i,method in enumerate(methods):
  for phase,off,marker in [('cal',-.13,'o'),('audit',.13,'s')]:
   x=next(x for x in B if x['dataset']==ds and x['phase']==phase and x['method']==method and x['reference']=='base');v=x['delta_recall_pp'];lo=x['recall_CI_low_pp'];hi=x['recall_CI_high_pp']
   ax.errorbar(i+off,v,yerr=[[v-lo],[hi-v]],fmt=marker,c=colors[method],mfc=colors[method] if phase=='audit' else 'white',capsize=3,ms=5)
 ax.axhline(0,c='#666666',lw=.9);ax.set_xticks(range(5),[labels[x].replace(' ','\n') for x in methods]);ax.set_ylabel('Recall change vs original (pp)');ax.set_title(f'{names[ds]}: cal / audit (512 each)');ax.grid(axis='y',alpha=.15)
 ax=axes[1,col]
 for method,off in [('base',(7,8)),('rl_20260918',(7,-17)),('exact_20260918',(7,8))]:
  p=next(x for x in P if x['dataset']==ds and x['phase']=='wall' and x['method']==method);r=p['recall_pct'];v=p['QPS_median']
  ax.errorbar(r,v,yerr=[[v-p['QPS_min']],[p['QPS_max']-v]],fmt='o',c=colors[method],capsize=4,ms=7);ax.annotate(labels[method],(r,v),xytext=off,textcoords='offset points',fontsize=9,c=colors[method])
 ax.set_title(f'{names[ds]}: actual L128 wall');ax.set_xlabel('Audit Recall@10 (%) — zoomed scale');ax.set_ylabel('Full-path QPS — zoomed scale');ax.margins(x=.3,y=.3);ax.grid(alpha=.15)
fig.suptitle('Development quality and measured inference cost',fontsize=16)
save(fig,'development_quality_and_cost','Top: hollow circle = cal, filled square = audit; paired query-bootstrap 95% intervals (5000 resamples).\nBottom: median and min–max of 5 serial single-thread plain-wall runs. One L, historical queries; not a frontier.')
outcomes=[];steps=[];actions=[]
for ds in DS:
 for phase in ['cal','audit']:
  for method in methods:
   qq=[x for x in Q if x['dataset']==ds and x['phase']==phase and x['method']==method];p=next(x for x in P if x['dataset']==ds and x['phase']==phase and x['method']==method)
   outcomes.append(dict(dataset=ds,phase=phase,method=method,helped=p['helped'],harmed=p['harmed'],same=512-p['helped']-p['harmed'],changed=p['changed_queries'],first_stop_pct=p['first_stop_pct']))
   for s in range(4):steps.append(dict(dataset=ds,phase=phase,method=method,swaps=s,queries=sum(x['swaps']==s for x in qq)))
   a=np.fromfile(O/'results'/f'{phase}_{ds}'/method/'actions.i32',np.int32).reshape(512,3)
   for t in range(3):
    for act in range(-1,16):actions.append(dict(dataset=ds,phase=phase,method=method,step=t,action=act,queries=int((a[:,t]==act).sum())))
for name,data in [('swap_outcomes',outcomes),('swap_counts',steps),('action_counts',actions)]:
 dump(O/'results'/f'{name}.json',data)
 with (O/'results'/f'{name}.csv').open('w') as f:w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
fig,axes=plt.subplots(2,2,figsize=(12,8))
for col,ds in enumerate(DS):
 a=[x for x in outcomes if x['dataset']==ds and x['phase']=='audit'];ax=axes[0,col];bottom=np.zeros(5)
 for key,color in [('helped','#087f9c'),('same','#e1e6e9'),('harmed','#c15a3b')]:
  v=np.array([x[key] for x in a]);ax.bar(range(5),v,bottom=bottom,color=color,label=key);bottom+=v
 ax.set_xticks(range(5),[labels[x].replace(' ','\n') for x in methods]);ax.set_ylabel('Queries (of 512)');ax.set_title(f'{names[ds]}: audit hit-count outcomes');ax.legend(fontsize=8,ncol=3,loc='upper center',bbox_to_anchor=(.5,1.02));ax.set_ylim(0,560)
 ax=axes[1,col];bottom=np.zeros(5)
 for s,color in enumerate(['#dbe4eb','#91c6ce','#3c9aaa','#135368']):
  v=np.array([next(x['queries'] for x in steps if x['dataset']==ds and x['phase']=='audit' and x['method']==m and x['swaps']==s) for m in methods]);ax.bar(range(5),v,bottom=bottom,color=color,label=f'{s} swaps');bottom+=v
 ax.set_xticks(range(5),[labels[x].replace(' ','\n') for x in methods]);ax.set_ylabel('Queries (of 512)');ax.set_title(f'{names[ds]}: executed swap counts');ax.legend(fontsize=8,ncol=4,loc='upper center',bbox_to_anchor=(.5,1.02));ax.set_ylim(0,560)
fig.suptitle('Changing entries is not sufficient for improving search',fontsize=16)
save(fig,'swap_behavior','All methods use the same top8 candidate pool, action domain and three-step limit. Seed18 was selected on cal.\nEqual hit count may still involve different entries or returned IDs. Stop does not guarantee nonnegative true reward.')
wb=Workbook();wb.remove(wb.active)
files=['models','training_checkpoints','training_curves','points','paired','timings','query_results','swap_outcomes','swap_counts','action_counts']
sources=[]
for name in files:
 ws=wb.create_sheet(name[:31]);p=O/'results'/f'{name}.csv'
 with p.open() as f:
  rr=list(csv.reader(f))
 for row in rr:
  cells=[]
  for value in row:
   if value in ['True','False']:cells.append(value=='True')
   elif not value:cells.append(None)
   else:
    try:cells.append(int(value))
    except ValueError:
     try:cells.append(float(value))
     except ValueError:cells.append(value)
  ws.append(cells)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174E64');c.alignment=Alignment(wrap_text=True)
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=min(38,max(13,len(str(col[0].value))+2))
 ws.row_dimensions[1].height=32
 sources.append(dict(sheet=ws.title,path=str(p.relative_to(O)),rows=len(rr)-1,sha256=sha(p)))
ws=wb.create_sheet('Sources');ws.append(['sheet','CSV relative path','rows','CSV SHA256'])
for x in sources:ws.append(list(x.values()))
ws.append(['Notes','All splits historically used; one L wall is a cost diagnostic. CI is descriptive, not a method GO.'])
ws.append(['Nested data','See paired.json for utility intervals and points.json for raw timing lists.'])
ws.freeze_panes='A2';ws.column_dimensions['A'].width=25;ws.column_dimensions['B'].width=100;ws.column_dimensions['D'].width=70
wb.save(O/'results/query_swap_policy.xlsx');dump(O/'reviews/workbook_sources.json',sources)
resources();print('FIGURES_AND_WORKBOOK_COMPLETE',flush=True)
