from common import *
import sys,csv,math,re
from urllib.parse import unquote
sys.path.insert(0,str(R/'entry_method_exploration/20260911_gt10_shared_transform_v1/src/vendor'))
from openpyxl import load_workbook

assert read(O/'reviews/training_verification.json')['status']=='PASS'
assert read(O/'reviews/output_verification.json')['status']=='PASS'
assert read(O/'reviews/independent_fit_deployment.json')['status']=='PASS'
assert read(O/'reviews/visual_review.json')['status']=='PASS'
assert read(O/'results/deployment_complete.json')['frontier_points']==50
check(O/'reviews/upstream.json');check(O/'reviews/before_training.json');check(O/'reviews/models_after_training.json');check(O/'reviews/E36_preserved.json')
if (O/'reviews/before_analysis.json').exists():check(O/'reviews/before_analysis.json')
old_count=0
for x in read(O/'reviews/old_artifacts_stat.json'):
 s=Path(x['path']).stat();assert s.st_size==x['bytes'] and s.st_mtime_ns==x['mtime_ns'],x['path'];old_count+=1
jobs=read(O/'logs/jobs.json');assert len(jobs)==18 and all(j['status']=='COMPLETE' and j['actual']<=j['reserved'] for j in jobs)
actual=sum(j['actual'] for j in jobs);reserved=sum(j['reserved'] for j in jobs);assert actual<=reserved<=CAP
assert actual==2211520,actual
ordered=sorted(jobs,key=lambda j:j['started_at'])
assert all(a['closed_at']<=b['started_at'] for a,b in zip(ordered,ordered[1:])), 'overlapping measured jobs'
training=read(O/'reviews/training_verification.json');assert training['optimizer_updates_verified']==1600 and training['action_W_hashes_verified']==12800 and training['checkpoints_verified']==88
processes=read(O/'logs/training_processes.json');assert len(processes)==2
assert all(p['status']=='COMPLETE' and p['peak_threads']==1 and p['peak_RSS_KiB']<32*(1<<20) for p in processes)
builds=sum(read(O/'results'/ds/'native_accounting.json')['builds'] for ds in DS)+sum(x['builds'] for x in read(O/'reviews/native_preflight.json'))
assert builds==12912 and builds<=15000
selected=read(O/'results/cal_selection.json')['selection'];check(O/'reviews/selection_before_audit.json')
for ds in DS:
 for arm in ARMS:
  records=read(O/'results'/ds/'cal_selection_details.json');best=max([x for x in records if x['model'].startswith(arm+'_')],key=lambda x:(x['mean_U'],-x['total_NDC'],-x['seed']))
  assert selected[ds][arm]==best
 for phase in ['fit_parity','cal','audit','frontier']:
  before=O/'reviews'/f'{phase}_{ds}_before.json';check(before)
  c=read(O/'inputs/jobs'/f'{phase}_{ds}.json');assert c['threads']==1 and c['query_cpu']==15
  for key in ['queries_file','gt_file','row_ids']:assert Path(c[key]).is_relative_to(O/'inputs'/ds)
  out=read(O/'results'/f'{phase}_{ds}'/'results.json');tc=read(O/'results'/f'{phase}_{ds}'/'thread_contract.json')
  assert tc['process_threads']==1 and tc['actual_openmp_threads']==1 and out['GT_opened_after_all_raw_search']
  assert out['timing_repeats']==(5 if phase=='frontier' else 0)
  assert all(m['full_plain_count_parity_queries']==c['num_queries'] and m['original_kernel_parity_queries']==16 and m['repeated_output_parity'] for m in out['methods'])
assert sha(O/'bin/evaluate')==sha(E34/'bin/evaluate')

workbook=load_workbook(O/'results/offline_profile_rl.xlsx',read_only=True,data_only=False);cells=0;csv_rows=0
for name in workbook.sheetnames:
 with (O/'results'/f'{name}.csv').open(newline='') as f:
  rows=list(csv.reader(f))
 values=workbook[name].iter_rows(values_only=True)
 for expected,current in zip(rows,values,strict=True):
  assert len(expected)==len(current)
  for a,b in zip(expected,current):
   if a=='':assert b is None
   elif isinstance(b,bool):assert a.lower()==str(b).lower()
   elif isinstance(b,(int,float)):assert math.isclose(float(a),b,rel_tol=1e-12,abs_tol=1e-10)
   else:assert a==str(b)
   cells+=1
  csv_rows+=1
workbook.close()
docs=[R/x for x in ['AGENTS.md','summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md','experiment_history/26_offline_profile_rl.md']]+[O/x for x in ['report.md','protocol.md','method_basis.md']]
dump(O/'reviews/final_audit.json',dict(status='AUDITING_DOCUMENT_LINKS'))
links=0
for doc in docs:
 text=doc.read_text()
 for href in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',text):
  if '://' in href or href.startswith('#'):continue
  path=unquote(href.split('#')[0].split(' "')[0]);target=(doc.parent/path).resolve();assert target.exists(),(doc,path);links+=1
assert 'E37' in (R/'summary.md').read_text() and 'E37' in (R/'s-target.md').read_text() and 'E37' in (R/'s-experiment.md').read_text()
resources();sizes=sum(p.stat().st_size for p in O.rglob('*') if p.is_file())
report=dict(status='PASS',training_runs=8,optimizer_updates=1600,sampled_static_profiles=12800,HNSW_builds=builds,new_NSG_calls=actual,reserved_NSG_calls=reserved,
 raw_training_and_checkpoint_outputs_set_scored=training['raw_query_outputs_set_scored'],checkpoints_verified=88,optimizer_and_action_replay='PASS',
 independent_fit_deployment_queries=1280,baseline_fit_parity_queries=10240,baseline_cal_audit_regression_queries=4096,
 deployment_points=80,frontier_points=50,independent_plain_wall_repeats=250,actual_query_threads=1,plain_count_original_kernel_and_cross_L_consistency='PASS',
 prior_E36_entire_SHA_checked=True,prior_artifacts_stat_checked=old_count,upstream_files_SHA_checked=len(read(O/'reviews/upstream.json')['files']),
 validation_heldout_raw_input_read_or_hash=False,historical_fit_cal_audit_exposure=True,default_replaced=False,independent_formal_GO=False,
 workbook_cells_checked=cells,CSV_rows_checked=csv_rows,documents_checked=len(docs),local_links_checked=links,visual_review='PASS',artifact_bytes_before_final_freeze=sizes,free_disk_bytes=shutil.disk_usage(O).free,closed_at=time.time())
dump(O/'reviews/final_audit.json',report)
freeze(O/'reviews/final_artifact_freeze.json',[p for p in O.rglob('*') if p.is_file() and p.name not in ['state.json','final_artifact_freeze.json'] and not p.name.endswith('.log')]+docs)
dump(O/'state.json',dict(status='COMPLETE_OFFLINE_PROFILE_RL',training_runs=8,optimizer_updates=1600,NSG_calls=actual,HNSW_builds=builds,development_gate=read(O/'results/development_gate.json')['status'],final_audit='PASS',default_replaced=False,validation_heldout_raw_input_access=False,closed_at=time.time()))
print('FINAL_AUDIT_PASS',actual,'NSG calls',builds,'HNSW builds',cells,'Excel cells',links,'links',flush=True)
