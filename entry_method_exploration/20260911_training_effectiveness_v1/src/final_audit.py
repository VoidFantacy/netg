from common import *
import sys,re,math,xml.etree.ElementTree as ET
sys.dont_write_bytecode=True;sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import load_workbook
import numpy as np

dump(O/'reviews/final_audit.json',dict(status='RUNNING',timestamp=time.time()))
assert read(O/'state.json')['status']=='REPORT_AND_DOCS_COMPLETE_AUDIT_PENDING'
required=['preflight','objective_tests','model_and_cache_verification','checkpoint_diagnostics','plots_generated','plot_visual_review','report_review','document_update']
for name in required:
    review=read(O/'reviews'/(name+'.json'));assert review['status']=='PASS',name
    for item in review.get('files',[]):assert identity(item['path'])==item,item['path']
protected=read(O/'reviews/protected_before.json')
for item in protected:assert identity(item['path'])==item,item['path']
checks=0
for path in [O/'reviews/frozen_before_training.json',O/'models/frozen_after_training.json',O/'reviews/frozen_analysis_before_final_audit.json']:
    check_freeze(path);checks+=len(read(path)['files'])
assert not any('heldout' in x['path'].lower() or 'held_out' in x['path'].lower() for x in protected)
assert read(O/'logs/training_process.json')['wall_seconds']<=600
budget=read(O/'models/training_budget.json');assert budget['status']=='PASS' and budget['all_runs_complete'] and budget['updates']==4800 and budget['runs']==4
assert budget['elapsed_seconds']<=600
runs=read(O/'models/training_runs.json');assert len(runs)==4
for r in runs:
    assert r['status']=='TRAINED_FIT_ONLY' and not r['cal_used'] and not r['audit_used'] and not r['validation_used']
    curve=read(O/'models'/r['dataset']/r['cell']/'training_curve.json');assert len(curve)==1201
    best=float('inf');best_epoch=0
    for i,row in enumerate(curve):
        assert row['epoch']==i and abs(row['objective']-(row['data_loss']+100*row['drift']))<1e-12
        if row['objective']<best:best=row['objective'];best_epoch=i
        assert row['best_objective']==best and row['best_epoch']==best_epoch
        step=i+1
        lr=0 if i==1200 else 3e-5+(3e-4-3e-5)*(step-1)/19 if step<=20 else 3e-4 if step<=200 else 3e-6+(3e-4-3e-6)*(1+math.cos(math.pi*(step-200)/1000))/2
        assert abs(row['next_lr']-lr)<1e-16
    assert r['selected_epoch']==best_epoch
tables=read(O/'results/analysis_tables.json');assert len(tables)==6 and len(tables['checkpoint_sensitivity'])==4
assert len(tables['model_comparison'])==4 and len(tables['training_curve'])==5128 and len(tables['cache_metrics'])==30 and len(tables['paired_cache'])==42
assert len(tables['alignment_curve'])==528
mv=read(O/'reviews/model_and_cache_verification.json');assert len(mv['models'])==8 and mv['logical_labels_reused']==2240000 and mv['new_labels']==mv['NSG_calls']==mv['optimizer_steps']==0
gate=read(O/'results/cache_gate.json');assert gate['status']=='STOP_CACHE_UTILITY_GATE' and gate['passing_cells']==[]
for row in gate['by_dataset']:
    ds=row['dataset'];cell=row['cell'];pts={x['method']:x for x in tables['cache_metrics'] if x['dataset']==ds and x['group']=='all'}
    cur,old,base=pts['new_'+cell],pts['old_'+cell],pts['anchor']
    conditions=[cur['candidate_recall_pct']-base['candidate_recall_pct']>=.2-1e-10,cur['candidate_recall_pct']-old['candidate_recall_pct']>1e-10,
        cur['candidate_graph_ndc']/base['candidate_graph_ndc']-1<=.01,cur['full_exact_known_fraction']>=.995,
        cur['full_exact_recall_lower_pct']-base['full_exact_recall_upper_pct']>=.2-1e-10]
    assert conditions==list(row['conditions'].values()) and row['passed']==all(conditions)
for ds in DS:
    for name in ['anchor','old_single','old_multi','new_single','new_multi']:
        z=np.load(O/'results'/ds/(name+'.npz'));assert len(z['slot'])==20000
        assert np.all((z['candidate_hits']>=0)&(z['candidate_hits']<=10))
        for group in ['all','headroom','no_headroom']:
            p=next(x for x in tables['cache_metrics'] if x['dataset']==ds and x['method']==name and x['group']==group)
            assert p['full_exact_recall_lower_pct']<=p['full_exact_recall_upper_pct']
            assert abs(p['full_exact_recall_upper_pct']-p['full_exact_recall_lower_pct']-100*p['unknown_queries']/p['queries'])<1e-10
wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);cells=0
assert wb.sheetnames==list(tables)+['Sources']
for name,rows in tables.items():
    ws=wb[name];values=list(ws.values);keys=list(dict.fromkeys(k for r in rows for k in r))
    assert list(values[0])==keys and len(values)==len(rows)+1
    for values_row,row in zip(values[1:],rows):
        for v,k in zip(values_row,keys):
            ref=row.get(k)
            if isinstance(ref,float):assert v is not None and math.isclose(v,ref,rel_tol=1e-12,abs_tol=1e-12),(name,k,v,ref)
            else:assert v==ref,(name,k,v,ref)
            cells+=1
plots=read(O/'reviews/plots_generated.json')['plots'];assert len(plots)==3
for name in plots:
    png=O/'figures'/(name+'.png');svg=O/'figures'/(name+'.svg')
    assert png.read_bytes()[:8]==b'\x89PNG\r\n\x1a\n' and png.stat().st_size>20000
    ET.parse(svg)
docs=[O/'report.md',O/'protocol.md',O/'README.md',R/'summary.md',R/'spec.md',R/'s-target.md',R/'s-experiment.md',R/'experiment_history/README.md',R/'experiment_history/20_training_effectiveness.md']
links=0
for doc in docs:
    s=doc.read_text()
    for target in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',s):
        if target.startswith(('http:','https:','#','mailto:')):continue
        p=target.split('#')[0].strip('<>');path=Path(p) if p.startswith('/') else doc.parent/p
        assert path.exists(),(doc,p);links+=1
subprocess.run(['git','diff','--check','--','summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md'],cwd=R,check=True)
size=sum(p.stat().st_size for p in O.rglob('*') if p.is_file());free=shutil.disk_usage(O).free
assert size<=512*(1<<20) and free>=10*(1<<30)
dump(O/'reviews/final_audit.json',dict(status='PASS',protected_identity_checks=len(protected),frozen_identity_checks=checks,training_runs=4,optimizer_updates=4800,
    total_training_seconds=budget['elapsed_seconds'],new_labels=0,logical_labels_reused=2240000,real_search_calls=0,new_HNSW_queries=0,new_QPS_points=0,
    model_checkpoint_verifications=8,cache_models_scored=10,workbook_verified_cells=cells,figures_visually_reviewed=3,document_links_checked=links,
    artifact_bytes=size,free_bytes=free,heldout_access=False,timestamp=time.time()))
dump(O/'state.json',dict(status='COMPLETE',conclusion='TRAINING_EFFECTIVENESS_COMPLETE_OPTIMIZATION_IMPROVED_CACHE_GATE_FAILED',training_runs=4,
    real_search_calls=0,new_labels=0,new_QPS_points=0,heldout_access=False))
print('FINAL_AUDIT_PASS',len(protected),checks,cells,links,flush=True)
