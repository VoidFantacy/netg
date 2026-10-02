from common import *
import sys
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

SHORT={DS[0]:'WebVid2.5M',DS[1]:'Yandex10M IP'}
COLORS={'single':'#2375b7','multi':'#df7a23'}
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none'})

def save(fig,name):
    fig.savefig(O/'figures'/(name+'.png'),dpi=160,facecolor='white');fig.savefig(O/'figures'/(name+'.svg'),facecolor='white');plt.close(fig)

assert read(O/'state.json')['status']=='CACHE_ANALYSIS_COMPLETE'
t=read(O/'results/analysis_tables.json')
fig,axs=plt.subplots(2,2,figsize=(15,9.2))
for j,ds in enumerate(DS):
    allnew=[]
    for cell in CELLS:
        for exp,style in [('E30','--'),('E31','-')]:
            rows=[r for r in t['training_curve'] if r['dataset']==ds and r['cell']==cell and r['experiment']==exp]
            label=cell.capitalize()+(' / original 80 steps' if exp=='E30' else ' / bounded optimization')
            axs[0,j].plot([r['epoch'] for r in rows],[r['objective'] for r in rows],style,color=COLORS[cell],label=label,linewidth=1.6,alpha=1 if exp=='E31' else .6)
            if exp=='E31':allnew.extend(r['objective'] for r in rows)
            ar=[r for r in t['alignment_curve'] if r['dataset']==ds and r['cell']==cell and r['experiment']==exp and r['group']=='all' and r['teacher']=='multi']
            base=ar[0]['candidate_recall_pct']
            axs[1,j].plot([r['epoch'] for r in ar],[r['candidate_recall_pct']-base for r in ar],style,color=COLORS[cell],marker='o',markersize=3,label=label,linewidth=1.6,alpha=1 if exp=='E31' else .6)
        selected=next(r for r in t['model_comparison'] if r['dataset']==ds and r['cell']==cell)
        utility=next(r for r in t['paired_cache'] if r['dataset']==ds and r['method']=='new_'+cell and r['reference']=='anchor' and r['group']=='all')
        axs[1,j].scatter([selected['selected_epoch']],[utility['candidate_delta_recall_pp']],marker='D',s=35,color=COLORS[cell],edgecolor='#172334',linewidth=.6,zorder=6,label='Selected float32 export' if cell=='single' else None)
    lo,hi=min(allnew),max(allnew);pad=(hi-lo)*.12
    axs[0,j].set_ylim(lo-pad,hi+pad);axs[0,j].set_title(SHORT[ds]+' — unchanged training objective')
    axs[0,j].set_ylabel('Data loss + 100 × profile displacement')
    axs[0,j].text(.98,.96,'Original early overshoot exceeds this view;\nsee first_steps for the full values.',transform=axs[0,j].transAxes,ha='right',va='top',fontsize=9,color='#66727e')
    axs[1,j].axhline(.2,color='#339266',linestyle=':',linewidth=1.5,label='Cache gate: +0.20 pp')
    axs[1,j].axhline(0,color='#566473',linewidth=.8);axs[1,j].set_title('Selected C8 triple: real cached L128 hits')
    axs[1,j].set_ylabel('Recall@10 change from initialization (pp)')
    for ax in axs[:,j]:ax.set_xlabel('Optimizer updates');ax.grid(alpha=.15)
fig.suptitle('Better optimization versus actual labeled search utility',fontsize=18,y=.98)
hs,ls=axs[1,0].get_legend_handles_labels();fig.legend(hs,ls,loc='lower center',bbox_to_anchor=(.5,.045),ncol=3,frameon=False,fontsize=10)
fig.text(.5,.015,'Lines: float64 training trajectory; diamonds: selected float32 exports. Cached L128 utility, without a new HNSW / NSG run.',ha='center',fontsize=9,color='#566473')
fig.subplots_adjust(left=.075,right=.98,top=.91,bottom=.19,hspace=.37,wspace=.23);save(fig,'optimization_and_utility')

fig,axs=plt.subplots(2,2,figsize=(14,8.4))
for j,ds in enumerate(DS):
    for i,cell in enumerate(CELLS):
        ax=axs[i,j]
        for exp,color in [('E30','#9a7373'),('E31',COLORS[cell])]:
            rows=[r for r in t['training_curve'] if r['dataset']==ds and r['cell']==cell and r['experiment']==exp and r['epoch']<=20]
            ax.plot([r['epoch'] for r in rows],[r['objective'] for r in rows],marker='o',markersize=3,color=color,label=exp,linewidth=1.5)
        ax.set_title(SHORT[ds]+' / '+cell);ax.set_xlabel('Optimizer updates');ax.set_ylabel('Full training objective');ax.grid(alpha=.15)
fig.suptitle('The first 20 steps: original versus fixed smaller-step schedule',fontsize=17,y=.98)
fig.legend(*axs[0,0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,.035),ncol=2,frameon=False)
fig.text(.5,.015,'An increasing objective alone does not establish an implementation bug. No learning-rate grid was run.',ha='center',fontsize=9,color='#566473')
fig.subplots_adjust(left=.075,right=.98,top=.90,bottom=.16,hspace=.42,wspace=.22);save(fig,'first_steps')

fig,axs=plt.subplots(1,2,figsize=(14,5.8))
names=['old_single','old_multi','new_single','new_multi'];labels=['Old single','Old multi','New single','New multi']
for ax,ds in zip(axs,DS):
    rows=[next(r for r in t['paired_cache'] if r['dataset']==ds and r['method']==name and r['reference']=='anchor' and r['group']=='all') for name in names]
    ys=np.array([r['candidate_delta_recall_pp'] for r in rows]);lo=np.array([r['ci95_low_pp'] for r in rows]);hi=np.array([r['ci95_high_pp'] for r in rows])
    ax.bar(range(4),ys,color=['#90b4d1','#e4be93',COLORS['single'],COLORS['multi']],width=.62)
    ax.errorbar(range(4),ys,yerr=[ys-lo,hi-ys],fmt='none',ecolor='#22364a',capsize=4)
    ax.axhline(.2,color='#339266',linestyle=':',label='Cache gate: +0.20 pp');ax.axhline(0,color='#566473',linewidth=.8)
    ax.set_xticks(range(4),labels);ax.set_title(SHORT[ds]);ax.set_ylabel('Cached Recall change vs anchor (pp)');ax.grid(axis='y',alpha=.15)
fig.suptitle('Actual label utility at the selected training checkpoint',fontsize=17,y=.98)
fig.text(.5,.055,'Whiskers: paired fit-query bootstrap 95% intervals (2,000 draws); no multiplicity correction or fresh blind evaluation.',ha='center',fontsize=9,color='#566473')
fig.text(.5,.02,'The stage gate also checks cost, improvement over the old model, and worst-case full-H unknown-combination bounds.',ha='center',fontsize=9,color='#566473')
fig.subplots_adjust(left=.075,right=.98,top=.87,bottom=.20,wspace=.22);save(fig,'cache_utility')

wb=Workbook();wb.remove(wb.active)
for name,rows in t.items():
    ws=wb.create_sheet(name);fields=list(dict.fromkeys(k for r in rows for k in r));ws.append(fields)
    for row in rows:ws.append([row.get(k) for k in fields])
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23384D')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(14,len(str(col[0].value))+2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,float):cell.number_format='0.000000'
ws=wb.create_sheet('Sources');ws.append(['Item','Definition / file'])
for item,value in [('Protocol',str(O/'protocol.md')),('Old models',str(E30/'models')),('Raw labels',str(E28/'labels')),
    ('Current report',str(O/'report.md')),('Objective','Same E30 normalized shared transform, lambda100, C8 set hinge; only optimization changed'),
    ('Utility','Previously measured real L128/K10 labels for each of 56 C8 triples; no new QPS'),
    ('Full-H bound','Unknown triple hits range from 0 to 10; conditional means do not replace full-query bounds'),
    ('Intervals','2,000 paired-query bootstrap draws over empirical integer hit differences; exploratory fit-only')]:ws.append([item,value])
ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=115
wb.save(O/'chart_data.xlsx')
dump(O/'reviews/plots_generated.json',dict(status='PASS',plots=['optimization_and_utility','first_steps','cache_utility'],tables=list(t)))
print('PLOTS_AND_WORKBOOK_COMPLETE',flush=True)
