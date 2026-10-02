from analyze import *
import sys
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker

COL={'gt10_hard':'#25364b','cost_ce':'#2b8cbe','keep_ce':'#e08c26',
     'cost_margin':'#087f8c','keep_margin':'#d55439','nsg':'#969da6',
     'fast30k':'#8b8c39','gt1_hard':'#9c7ac3','gt1_soft':'#c6b4d9',
     'gt10_soft':'#896442','e28_mlp':'#cb65a0'}
LINE={'cost_ce':'--','keep_ce':'--','cost_margin':'-','keep_margin':'-','gt10_hard':'-'}
MARK={'cost_ce':'o','keep_ce':'s','cost_margin':'^','keep_margin':'D','gt10_hard':'o'}
plt.rcParams.update({'font.size':10,'axes.titlesize':13,'axes.labelsize':11,'figure.dpi':150,
                     'savefig.dpi':180,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False})

def save(fig,name):
    for ext in ['png','svg']:
        fig.savefig(O/'figures'/f'{name}.{ext}',bbox_inches='tight',facecolor='white')
    plt.close(fig)

def curve_figure(rows,methods,name,title):
    fig,axes=plt.subplots(1,2,figsize=(13.5,5.7))
    handles=[]
    for ax,ds in zip(axes,DS):
        for method in methods:
            ms=sorted([r for r in rows if r['dataset']==SHORT[ds] and r['stage']=='validation' and r['method']==method],key=lambda r:r['L'])
            x=np.array([r['recall_pct'] for r in ms]);y=np.array([r['qps'] for r in ms])
            err=np.array([[r['qps']-r['qps_min'] for r in ms],[r['qps_max']-r['qps'] for r in ms]])
            artist=ax.errorbar(x,y,yerr=err,label=NAMES[method],color=COL[method],linestyle=LINE.get(method,':'),
                        marker=MARK.get(method,'.'),markersize=4,linewidth=1.8 if method in CELLS+['gt10_hard'] else 1.1,
                        capsize=2,alpha=1 if method in CELLS+['gt10_hard'] else .8)
            if ax==axes[0]:handles.append(artist)
        ax.set(title=SHORT[ds],xlabel='Recall@10 (%)',ylabel='QPS (one thread; log scale)')
        ax.set_yscale('log');ax.yaxis.set_major_locator(ticker.FixedLocator([400,800,1000,2000,4000]))
        ax.yaxis.set_major_formatter(ticker.StrMethodFormatter('{x:.0f}'));ax.yaxis.set_minor_formatter(ticker.NullFormatter())
        ax.grid(alpha=.2,which='major')
        ax.text(.03,.035,'L = 64 / 128 / 256 / 400 / 800',transform=ax.transAxes,fontsize=9,color='#4b5563')
    fig.suptitle(title,y=.97,fontsize=15)
    ncol=3 if len(methods)<=5 else 4
    fig.legend(handles,[NAMES[m] for m in methods],loc='lower center',bbox_to_anchor=(.5,.025),ncol=ncol,frameon=False,fontsize=9)
    fig.subplots_adjust(left=.065,right=.985,top=.85,bottom=.28,wspace=.25)
    save(fig,name)

def interpolation_figure(rows):
    fig,axes=plt.subplots(1,2,figsize=(12.5,4.9))
    handles=[]
    for ax,ds in zip(axes,DS):
        for cell in CELLS:
            ms=sorted([r for r in rows if r['dataset']==SHORT[ds] and r['method']==cell and r['predeclared_target']],key=lambda r:r['recall_target_pct'])
            x=[r['recall_target_pct'] for r in ms];y=[r['delta_qps_pct'] for r in ms]
            line,=ax.plot(x,y,color=COL[cell],linestyle=LINE[cell],marker=MARK[cell],label=NAMES[cell])
            ax.fill_between(x,[(r['ratio_min']-1)*100 for r in ms],[(r['ratio_max']-1)*100 for r in ms],color=COL[cell],alpha=.1)
            if ax==axes[0]:handles.append(line)
        ax.axhline(0,color='#364152',linewidth=.8);ax.axhline(3,color='#6d9665',linewidth=.7,linestyle=':')
        ax.set(title=SHORT[ds],xlabel='Recall target (%)',ylabel='Interpolated QPS change vs GT10 (%)')
        ax.set_xticks([v*100 for v in TARGETS[ds]]);ax.grid(alpha=.2)
    fig.suptitle('Same-recall comparison — INTERPOLATION, no extrapolation',y=.98,fontsize=14)
    fig.legend(handles,[NAMES[m] for m in CELLS],loc='lower center',ncol=2,frameon=False,bbox_to_anchor=(.5,.005))
    fig.subplots_adjust(left=.075,right=.985,top=.86,bottom=.25,wspace=.25)
    save(fig,'interpolated_qps')

def factor_figure(rows):
    fig,axes=plt.subplots(3,2,figsize=(12.5,10))
    panels=[('Teacher: keep minus cost',['teacher_under_ce','teacher_under_margin']),
            ('Loss: margin minus CE',['loss_under_cost','loss_under_keep']),
            ('Interaction: difference of differences',['interaction'])]
    descriptions={'teacher_under_ce':'Under CE','teacher_under_margin':'Under margin',
                  'loss_under_cost':'Under cost teacher','loss_under_keep':'Under keep teacher','interaction':'Teacher × loss'}
    for row,(title,keys) in enumerate(panels):
        for col,ds in enumerate(DS):
            ax=axes[row,col]
            for j,key in enumerate(keys):
                ms=sorted([r for r in rows if r['dataset']==SHORT[ds] and r['stage']=='validation' and r['contrast']==key],key=lambda r:r['L'])
                color=['#147d92','#cf7138'][j]
                x=[r['L'] for r in ms];y=[r['delta_recall_pp'] for r in ms]
                ax.plot(x,y,marker='o',color=color,label=descriptions[key])
                ax.fill_between(x,[r['ci_low_pp'] for r in ms],[r['ci_high_pp'] for r in ms],color=color,alpha=.12)
            ax.axhline(0,color='#4b5563',linewidth=.8)
            ax.set_xscale('log',base=2);ax.set_xticks([64,128,256,400,800]);ax.xaxis.set_major_formatter(ticker.StrMethodFormatter('{x:.0f}'))
            ax.set(title=f'{SHORT[ds]} — {title}',xlabel='NSG search L',ylabel='Recall difference (pp)')
            ax.grid(alpha=.2);ax.legend(frameon=False,fontsize=9)
    fig.suptitle('2 × 2 controls: paired-query Recall effects at the same L',y=.995,fontsize=15)
    fig.tight_layout(rect=[0,0,1,.97]);save(fig,'factor_effects')

def alignment_figure(rows):
    fig,axes=plt.subplots(2,2,figsize=(12.5,7.8))
    handles=[]
    for col,ds in enumerate(DS):
        for row,teacher in enumerate(['cost','keep']):
            ax=axes[row,col]
            for cell in CELLS:
                ms=sorted([r for r in rows if r['dataset']==SHORT[ds] and r['cell']==cell and r['reference_teacher']==teacher and r['group']=='all'],key=lambda r:r['epoch'])
                line,=ax.plot([r['epoch'] for r in ms],[r['entry_overlap']*100 for r in ms],color=COL[cell],linestyle=LINE[cell],label=NAMES[cell])
                if row==0 and col==0:handles.append(line)
            ax.axhline(ms[0]['entry_overlap']*100,color='#64748b',linestyle=':',linewidth=.8)
            ax.set(title=f'{SHORT[ds]} — fixed {teacher} teacher',xlabel='Epoch',ylabel='Teacher-entry overlap (%)')
            ax.grid(alpha=.2)
    fig.suptitle('Exact full-H top3 alignment on fit queries',y=.99,fontsize=15)
    fig.legend(handles,[NAMES[m] for m in CELLS],loc='lower center',ncol=2,frameon=False,bbox_to_anchor=(.5,.0))
    fig.subplots_adjust(left=.075,right=.985,top=.89,bottom=.15,wspace=.24,hspace=.42)
    save(fig,'training_alignment')

def generalization_figure(rows):
    fig,axes=plt.subplots(1,2,figsize=(12.5,5))
    stages=['fit','cal','audit','validation'];handles=[]
    for ax,ds in zip(axes,DS):
        for cell in CELLS:
            ms=[next(r for r in rows if r['dataset']==SHORT[ds] and r['method']==cell and r['stage']==stage and r['L']==128 and not r['exact']) for stage in stages]
            y=[r['delta_recall_pp'] for r in ms]
            line,=ax.plot(range(4),y,color=COL[cell],linestyle=LINE[cell],marker=MARK[cell],label=NAMES[cell])
            if ax==axes[0]:handles.append(line)
        ax.axhline(0,color='#4b5563',linewidth=.8)
        ax.set(title=SHORT[ds],ylabel='Recall difference vs GT10 (pp)',xlabel='Query split')
        ax.set_xticks(range(4),['Fit 20K','Cal 1024','Audit 1024','Validation 2K']);ax.grid(alpha=.2)
    fig.suptitle('Actual HNSW → NSG Recall at L128, relative within each split',y=.98,fontsize=14)
    fig.legend(handles,[NAMES[m] for m in CELLS],loc='lower center',ncol=2,frameon=False,bbox_to_anchor=(.5,.0))
    fig.subplots_adjust(left=.075,right=.985,top=.85,bottom=.25,wspace=.25)
    save(fig,'fit_to_validation')

def workbook(tables):
    book=Workbook();book.remove(book.active)
    extra=[]
    for m in read(O/'models/training_runs.json'):
        extra.append(dict(dataset=SHORT[m['dataset']],cell=m['cell'],parameters=m['parameter_count'],
            training_seconds=m['training_seconds'],export_seconds=m['export_seconds'],
            initial_data_loss=m['first']['data_loss'],final_data_loss=m['last']['data_loss'],
            final_objective=m['last']['objective'],final_gradient_norm=m['last']['gradient_norm'],drift=m['last']['drift']))
    tables=dict(tables);tables['model_summary']=extra
    timing=[dict(dataset=r['dataset'],method=r['method'],L=r['L'],qps_median=r['qps'],qps_min=r['qps_min'],qps_max=r['qps_max'],
        range_over_median_pct=(r['qps_max']-r['qps_min'])/r['qps']*100) for r in tables['all_points'] if r['stage']=='validation']
    tables['timing_spread']=timing
    gd=read(O/'results/gradient_diagnostics.json')
    gr=[]
    for rec in gd['records']:
        for group,info in rec['groups'].items():
            gr.append(dict(dataset=SHORT[rec['dataset']],cell=rec['cell'],checkpoint=rec['checkpoint'],group=group,
                **info,gradient_cosine=rec['gradient_cosine'],projection_fraction=rec[group+'_projection_fraction']))
    tables['gradient_groups']=gr;tables['geometry']=gd['geometry']
    for name,rows in tables.items():
        if not rows:continue
        ws=book.create_sheet(name)
        fields=list(dict.fromkeys(k for row in rows for k in row))
        ws.append(fields)
        for row in rows:ws.append([row.get(k) for k in fields])
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='25364B')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            width=min(42,max(13,max(len(str(c.value or '')) for c in list(col)[:80])+1))
            ws.column_dimensions[col[0].column_letter].width=width
    ws=book.create_sheet('Sources');ws.append(['Item','Source','Notes'])
    ws.append(['Raw search outputs',str(O/'results'),'All 170 points, including 110 formal points.'])
    ws.append(['Protocol',str(O/'protocol.md'),'Frozen before model training and new queries.'])
    ws.append(['QPS','Three independent warm/plain batch timings','Single thread CPU15; timing repeats are not Recall samples.'])
    ws.append(['Interpolation','Recall versus log QPS on median Pareto points','No extrapolation; repeat ranges use the same median-frontier segment.'])
    ws.append(['Recall CI','2000 paired-query bootstrap draws','Seed20260911; no multiple-comparison correction.'])
    ws.append(['Historical datasets','fit/cal/audit/validation previously used','No fresh heldout generalization claim.'])
    for source in read(O/'sources/inventory.json'):
        ws.append([source['key'],source['url'],source['sha256']])
    ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=80;ws.column_dimensions['C'].width=75
    ws.freeze_panes='A2'
    book.save(O/'chart_data.xlsx')
    writecsv('model_summary',extra);writecsv('gradient_groups',gr);writecsv('geometry',gd['geometry'])
    writecsv('timing_spread',timing)

if __name__=='__main__':
    assert (O/'results/analysis_tables.json').exists()
    tables=read(O/'results/analysis_tables.json')
    curve_figure(tables['all_points'],['gt10_hard']+CELLS,'recall_qps_main','Teacher × objective alignment: actual Recall–QPS')
    curve_figure(tables['all_points'],list(NAMES),'recall_qps_all','Complete baseline and 2 × 2 comparison')
    interpolation_figure(tables['interpolation'])
    factor_figure(tables['factor_contrasts'])
    alignment_figure(tables['alignment_curves'])
    generalization_figure(tables['differences'])
    workbook(tables)
    print('6 PNG/SVG figures and chart_data.xlsx written',flush=True)
