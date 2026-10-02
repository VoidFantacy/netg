from common import *
import sys,math,ast,re,csv
from urllib.parse import unquote
sys.dont_write_bytecode=True;sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import load_workbook

check_only='--check-only' in sys.argv
state=read(O/'state.json');assert state['status']=='COMPLETE_DEVELOPMENT_STOP' and state['conclusion']=='HNSW_EDGE_RL_COMPLETE_WEBVID_SIGNAL_NO_DUAL_GAIN'
assert state['QPS_points']==0 and state['optimizer_updates']==2400 and state['training_runs']==4
resources()
old=read(O/'reviews/older_protection_reference.json')
for x in old:
 p=Path(x['path']);s=p.stat();assert s.st_size==x['bytes'] and s.st_mtime_ns==x['mtime_ns'],p
old_validation=[x for x in old if '/validation/' in x['path'] and Path(x['path']).name in ['queries.f32','gt.i64','row_ids.i64']]
freeze_paths=[O/'reviews/protected_before.json',O/'reviews/protocol_and_implementation_before_queries.json',O/'reviews/frozen_structural_repair.json',O/'reviews/initial_graphs.json',
 O/'reviews/models_after_training.json',O/'reviews/deployment_implementation.json',O/'reviews/selection_after_cal.json',O/'reviews/frozen_analysis_before_final_audit.json']
freeze_paths += [O/'reviews'/f'frozen_{split}_{ds}.json' for ds in DS for split in ['cal','audit']]
unique={}
for path in freeze_paths:
 for item in read(path)['files']:
  if item['path'] in unique:assert unique[item['path']]==item,item['path']
  unique[item['path']]=item
for name,item in unique.items():assert identity(name)==item,name
assert not any((O/'results').glob('validation_*')) and not any((O/'inputs').rglob('eval_validation.json'))
scope=read(O/'reviews/access_scope.json');assert scope['fit_decoded'] and scope['cal_decoded'] and scope['audit_decoded']
assert not scope['validation_decoded'] and not scope.get('validation_hash',False) and not scope['heldout_hash'] and not scope['heldout_decoded']
assert all('validation' not in str(x['path']).split('/')[-2:] for x in read(O/'reviews/protected_before.json')['files'] if Path(x['path']).name in ['queries.f32','gt.i64','row_ids.i64'])

jobs=read(O/'logs/jobs.json');assert len(jobs)==12 and all(j['status']=='COMPLETE' for j in jobs)
calls=sum(j['actual_calls'] for j in jobs);reserved=sum(j['reserved_calls'] for j in jobs);assert calls==500736 and reserved==500992 and reserved<CAP and state['NSG_calls']==calls
training=read(O/'logs/training_process.json');assert training['status']=='COMPLETE' and training['exit_code']==0 and training['process_threads_max']==1 and training['rss_peak_kib']<40*(1<<20)
for j in jobs:
 assert j['actual_calls']<=j['reserved_calls']
 if j['name'].startswith('train_'):assert j['updates']==600 and j['actual_calls']==104800 and j['training_seconds']<600
 if j['name'].startswith(('cal_','audit_')):assert j['process_threads_max']==1 and j['actual_calls']==10320
for pid in [training['pid']]+[j['pid'] for j in jobs if 'pid' in j]:
 p=Path('/proc')/str(pid)/'cmdline'
 if p.exists():assert str(O).encode() not in p.read_bytes(),'experiment process still alive'
assert sum(j.get('seconds',0) for j in jobs)<3600

train_review=read(O/'reviews/training_independent_verification.json');assert train_review['status']=='PASS' and train_review['optimizer_updates']==2400
assert sum(x['trajectories'] for x in train_review['datasets'])==307200 and sum(x['graph_artifacts_checked'] for x in train_review['datasets'])==36
assert max(x['max_numpy_replay_error'] for x in train_review['datasets'])<2.4e-15
assert read(O/'reviews/reinforce_tests.json')['status']=='PASS' and read(O/'reviews/router_preflight_v2.json')['status']=='PASS'
assert read(O/'reviews/first_preflight_failure.json')['NSG_calls']==read(O/'reviews/first_preflight_failure.json')['HNSW_query_calls']==0
source=read(O/'bin/evaluate.build.json')['argv'];assert '-I'+str(O/'src') not in source
assert '-I'+str(O/'src') in read(O/'bin/environment_v2.build.json')['argv']
for p in (O/'src').glob('*.py'):ast.parse(p.read_text())

# Historical root versions are retained before the authorized document updates.
original_roots={}
for path in [E32/'reviews/frozen_analysis_before_final_audit.json',E32/'reviews/frozen_after_access_scope_correction.json']:
 for item in read(path)['files']:
  p=Path(item['path'])
  if p.parent==R or p==R/'experiment_history/README.md':original_roots[str(p)]=item
backup_verified=0
for name,item in original_roots.items():
 p=O/'source_before'/Path(name).relative_to(R)
 if p.exists():
  q=identity(p);assert q['bytes']==item['bytes'] and q['mtime_ns']==item['mtime_ns'] and q['sha256']==item['sha256'];backup_verified+=1
assert backup_verified>=4
assert sha(O/'source_before/AGENTS.md')==sha(R/'AGENTS.md')

# Add a bounds check and baseline distance-bit regression to the already completed full GT/optimizer audits.
bounded_outputs=0;baseline_bit_outputs=0
for ds in DS:
 cfg=read(O/'inputs'/ds/'config.json');base_count=read(E24/'inputs'/ds/'prepared.json')['base_rows'];fit_rows=np.fromfile(cfg['row_ids'],np.int64)
 base=np.load(O/'results'/ds/'baseline.npz');oldbase=E32/'results'/('fit_v2_'+ds)/(policy(ds)['model_name']+'_L128')
 assert np.array_equal(base['distances'].reshape(-1).view(np.uint32),np.fromfile(oldbase/'distances.f32',np.float32).view(np.uint32));baseline_bit_outputs+=20000
 raw=[O/'results'/ds/'baseline.npz']+list((O/'models'/ds).glob('*/checkpoint_*/monitor.npz'))+list((O/'models'/ds).glob('*/trajectories/*.npz'))
 for p in raw:
  with np.load(p) as a:
   ids=a['ids'];assert (ids>=0).all() and (ids<base_count).all();bounded_outputs+=len(ids)
 development_rows=[]
 for split in ['cal','audit']:
  dest=O/'results'/f'{split}_{ds}';c=read(dest/'config.json');rs=read(dest/'results.json');v=read(dest/'independent_verification.json');th=read(dest/'thread_contract.json')
  assert v['status']=='PASS' and len(v['points'])==5 and rs['timing_repeats']==0 and rs['threads']==1 and th['process_threads']==th['actual_openmp_threads']==1 and th['cpu']==15
  rr=np.fromfile(c['row_ids'],np.int64);assert not np.intersect1d(fit_rows,rr).size;development_rows.append(rr)
  prev=E32/'results'/f'{split}_{ds}'/(policy(ds)['model_name']+'_L128')
  assert np.array_equal(np.fromfile(dest/'baseline_L128/distances.f32',np.float32).view(np.uint32),np.fromfile(prev/'distances.f32',np.float32).view(np.uint32));baseline_bit_outputs+=1024
  for m in rs['methods']:
   assert m['qps'] is None and not m['qps_runs'] and not m['end_to_end_seconds']
   ids=np.fromfile(dest/m['name']/'ids.i32',np.int32).reshape(1024,10);assert (ids>=0).all() and (ids<base_count).all();bounded_outputs+=1024
 assert not np.intersect1d(*development_rows).size
assert bounded_outputs==479680 and baseline_bit_outputs==44096

T=read(O/'results/report_tables.json');selection=read(O/'results/cal_selection.json')['selection'];gate=read(O/'results/development_gate.json')
assert gate['status']=='STOP_DEVELOPMENT_GATE'
for ds in DS:
 rs=[r for r in T['paired_development'] if r['dataset']==ds and r['method'].startswith('rl_') and r['split']=='cal']
 winner=max(rs,key=lambda r:(r['reward'],-int(r['method'].split('_')[1])));assert selection[ds]==int(winner['method'].split('_')[1])
 chosen=[r for r in T['paired_development'] if r['dataset']==ds and r['method']==f"rl_{selection[ds]}"]
 random=[r for r in T['paired_development'] if r['dataset']==ds and r['method']==f"random_{selection[ds]}"]
 passed=all((r['delta_pp']>=.05-1e-12 and r['total_change_pct']<=1+1e-12) or (r['delta_pp']>=-.02-1e-12 and r['total_change_pct']<=-1+1e-12) for r in chosen) and sum(r['reward'] for r in chosen)>sum(r['reward'] for r in random)
 assert next(x['pass_gate'] for x in gate['checks'] if x['dataset']==ds)==passed
 assert passed==(ds==DS[0])

# Verify every numeric/string cell in every table against the chart source; no plots used as data.
wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);cells=0
assert wb.sheetnames==list(T)+['Sources']
for name,rows in T.items():
 ws=wb[name];it=ws.iter_rows(values_only=True);keys=list(rows[0]);assert list(next(it))==keys;cells+=len(keys)
 actual=list(it);assert len(actual)==len(rows)
 for values,row in zip(actual,rows):
  for got,key in zip(values,keys):
   expected=row[key];cells+=1
   if isinstance(expected,float):assert isinstance(got,(int,float)) and math.isclose(got,expected,rel_tol=2e-14,abs_tol=1e-12),(name,key,got,expected)
   else:assert got==expected,(name,key,got,expected)
 with (O/'results/csv'/(name+'.csv')).open() as f:
  dr=csv.DictReader(f);assert dr.fieldnames==keys;rr=list(dr);assert len(rr)==len(rows)
  for a,b in zip(rr,rows):
   for key in keys:assert a[key]==('' if b[key] is None else str(b[key]))
for name,path,digest,scope in list(wb['Sources'].values)[1:]:assert Path(path).name==name and sha(path)==digest
wb.close()
assert read(O/'reviews/plot_visual_review.json')['status']=='PASS' and read(O/'reviews/report_review.json')['status']=='PASS'
docs=[R/'summary.md',R/'s-target.md',R/'s-experiment.md',R/'spec.md',R/'experiment_history/README.md',R/'experiment_history/22_hnsw_edge_rl.md',O/'report.md',O/'README.md',O/'protocol.md',O/'structural_preflight_repair.md',O/'literature/method_basis.md'];links=0
for p in docs:
 for target in re.findall(r'\[[^\]]+\]\(([^)]+)\)',p.read_text()):
  target=target.strip('<>').split('#')[0]
  if not target or target.startswith(('http://','https://','mailto:')):continue
  q=(p.parent/unquote(target)).resolve();assert q.exists(),(p,target);links+=1
assert 'E33 HNSW图边强化学习' in (R/'s-experiment.md').read_text() and '二十二份详细卡片' in (R/'s-experiment.md').read_text()
assert (R/'spec.md').read_text().count('## HNSW底层边强化学习 v1')==1
result=dict(status='PASS',conclusion=state['conclusion'],old_stat_protected_files=len(old),older_validation_stat_only_files=len(old_validation),
 SHA256_frozen_unique_files=len(unique),freeze_manifests=len(freeze_paths),root_backups_verified_against_E32=backup_verified,
 optimizer_updates_independently_replayed=2400,exported_graphs_verified=36,all_persisted_ID_outputs_GT_rescored_and_bounded=bounded_outputs,
 baseline_distance_bit_regression_outputs=baseline_bit_outputs,actual_NSG_calls=calls,reserved_NSG_calls=reserved,HNSW_queries=read(O/'results/costs.json')['HNSW_queries'],
 development_points=20,QPS_points=0,validation_hash=False,validation_decoded=False,heldout_hash=False,heldout_decoded=False,
 workbook_verified_cells=cells,workbook_sheets=len(T)+1,visually_reviewed_figures=3,local_document_links_checked=links,
 artifact_bytes=sum(p.stat().st_size for p in O.rglob('*') if p.is_file()),free_disk_bytes=shutil.disk_usage(O).free,no_new_queries_during_audit=True)
if not check_only:
 target=O/'reviews/final_audit.json'
 assert not target.exists() or read(target)['status']=='PENDING'
 dump(target,result)
print('FINAL_AUDIT_PASS',json.dumps(result),flush=True)
