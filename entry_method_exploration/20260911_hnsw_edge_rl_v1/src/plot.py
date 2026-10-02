from common import *
import csv
import sys
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
T=read(O/'results/report_tables.json');selection=read(O/'results/cal_selection.json')['selection']
N={DS[0]:'WebVid 2.5M',DS[1]:'Yandex 10M IP'};colors=['#147D92','#D76636'];plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':140,'savefig.dpi':180})
def save(fig,name):
 fig.savefig(O/'figures'/(name+'.png'),bbox_inches='tight');fig.savefig(O/'figures'/(name+'.svg'),bbox_inches='tight');plt.close(fig)

fig,axs=plt.subplots(2,2,figsize=(12,7),layout='constrained')
for i,ds in enumerate(DS):
 for j,seed in enumerate(SEEDS):
  rs=[r for r in T['training_curve'] if r['dataset']==ds and r['seed']==seed];y=np.array([r['reward'] for r in rs]);smooth=np.convolve(y,np.ones(50)/50,'valid')
  axs[i,0].plot(np.arange(50,601),smooth,label=f'Seed {seed%100}',color=colors[j],lw=1.6)
  ms=[r for r in T['monitor_checkpoints'] if r['dataset']==ds and r['seed']==seed];axs[i,1].plot([r['step'] for r in ms],[r['delta_pp'] for r in ms],'o-',label=f'Seed {seed%100}',color=colors[j],ms=4)
  chosen=next(r for r in T['selected_models'] if r['dataset']==ds and r['seed']==seed);axs[i,1].scatter(chosen['selected_step'],chosen['monitor_delta_pp'],marker='*',s=150,color=colors[j],edgecolor='white',linewidth=.6,zorder=5)
 for ax in axs[i]:ax.axhline(0,color='#667788',lw=.8,ls='--');ax.grid(alpha=.15);ax.set_xlabel('Optimizer update');ax.legend(loc='best')
 axs[i,0].set(title=N[ds]+' | sampled training episodes',ylabel='Reward (50-update rolling mean)')
 axs[i,1].set(title=N[ds]+' | exported static graphs',ylabel='Monitor Recall change (pp)')
fig.suptitle('Real search feedback trains HNSW edges; static-graph gains vary by seed',fontsize=13)
save(fig,'training_and_static_graphs')

fig,axs=plt.subplots(2,2,figsize=(12,8),layout='constrained')
for i,ds in enumerate(DS):
 rows=[r for r in T['paired_development'] if r['dataset']==ds];labels=[]
 for k,r in enumerate(rows):
  color=colors[0] if r['method'].startswith('rl_') else '#929BA7';seed=int(r['method'].split('_')[1]);split=r['split'];x=r['delta_pp'];lo=r['ci95_low_pp'];hi=r['ci95_high_pp']
  axs[0,i].errorbar(x,k,xerr=[[x-lo],[hi-x]],fmt='o' if seed==selection[ds] else 's',color=color,ms=5,capsize=2)
  labels.append(f"{split} / {'RL' if r['method'].startswith('rl_') else 'random'} {seed%100}"+(' *' if seed==selection[ds] else ''))
 axs[0,i].set_yticks(range(len(rows)),labels);axs[0,i].invert_yaxis();axs[0,i].axvline(0,color='#555',ls='--',lw=.8);axs[0,i].grid(axis='x',alpha=.15)
 axs[0,i].set(title=N[ds]+' | all seeds and matched random controls',xlabel='Recall change vs original graph (pp); paired 95% CI')
 pts=T['real_points'];selected=selection[ds];x=np.arange(2);width=.22
 for j,key in enumerate(['route_ndc','graph_ndc','total_ndc']):
  values=[]
  for split in ['cal','audit']:
   a=next(r for r in pts if r['dataset']==ds and r['split']==split and r['method']==f'rl_{selected}');b=next(r for r in pts if r['dataset']==ds and r['split']==split and r['method']=='baseline');values.append((a[key]/b[key]-1)*100)
  bars=axs[1,i].bar(x+(j-1)*width,values,width,label=['HNSW routing','NSG search','Total'][j],color=['#147D92','#F1B56A','#6554A4'][j])
  for bar,v in zip(bars,values):axs[1,i].text(bar.get_x()+width/2,v-.6 if v<0 else v+.4,f'{v:+.2f}%',ha='center',va='top' if v<0 else 'bottom',fontsize=8)
 axs[1,i].set_xticks(x,['cal','audit']);axs[1,i].set(title=f'{N[ds]} | cal-selected seed {selected%100}',ylabel='Distance computation change (%)');axs[1,i].axhline(0,color='#555',lw=.8);axs[1,i].set_ylim(-25,4);axs[1,i].legend(loc='lower right',fontsize=8)
fig.suptitle('Development: a WebVid signal, Yandex losses, and no measured QPS\n* Cal-selected seed; random controls match each node\'s retained degree',fontsize=13)
save(fig,'development_and_cost')

fig,axs=plt.subplots(1,2,figsize=(11,4.2),layout='constrained')
for i,ds in enumerate(DS):
 rows=[r for r in T['topology'] if r['dataset']==ds and r['seed']==selection[ds]];bins=np.arange(.5,33.5,1)
 axs[i].hist([r['original_outdegree'] for r in rows],bins=bins,histtype='step',lw=2,color='#929BA7',label='Original')
 axs[i].hist([r['learned_outdegree'] for r in rows],bins=bins,histtype='step',lw=2,color='#147D92',label='RL (random has identical degrees)')
 sm=next(r for r in T['selected_models'] if r['dataset']==ds and r['selected_by_cal'])
 axs[i].set(title=f"{N[ds]}: {sm['deleted_pct']:.1f}% of bottom-layer edges removed",xlabel='Outgoing degree',ylabel='Number of nodes');axs[i].legend(fontsize=8);axs[i].grid(alpha=.15)
fig.suptitle('The exported graphs change edges; coordinates, labels and upper layers are identical',fontsize=12)
save(fig,'graph_structure')

wb=Workbook();wb.remove(wb.active);csvdir=O/'results/csv';csvdir.mkdir(exist_ok=True)
for name,rows in T.items():
 if not rows:continue
 keys=list(rows[0]);ws=wb.create_sheet(name[:31]);ws.append(keys)
 with (csvdir/(name+'.csv')).open('w',newline='') as f:
  writer=csv.DictWriter(f,fieldnames=keys);writer.writeheader();writer.writerows(rows)
 for r in rows:ws.append([r[k] for k in keys])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17677A');cell.alignment=Alignment(wrap_text=True)
 for j,k in enumerate(keys,1):ws.column_dimensions[get_column_letter(j)].width=min(36,max(13,len(k)+2))
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   if isinstance(cell.value,float):cell.number_format='0.000000'
sources=wb.create_sheet('Sources');sources.append(['artifact','path','sha256','scope'])
for path,scope in [(O/'protocol.md','Prelocked method and gates'),(O/'structural_preflight_repair.md','Degree-one input compatibility correction'),(O/'results/report_tables.json','Exact source for every workbook table'),(O/'results/development_gate.json','No new validation or QPS'),(O/'reviews/training_independent_verification.json','Independent complete optimizer and raw-feedback verification')]:sources.append([path.name,str(path),sha(path),scope])
sources.freeze_panes='A2';sources.column_dimensions['A'].width=45;sources.column_dimensions['B'].width=95;sources.column_dimensions['C'].width=70;sources.column_dimensions['D'].width=70
wb.save(O/'chart_data.xlsx')
dump(O/'reviews/plots_generated.json',dict(status='GENERATED_PENDING_VISUAL_REVIEW',figures=['training_and_static_graphs','development_and_cost','graph_structure'],tables=list(T),workbook=str(O/'chart_data.xlsx')))
print('PLOTS_WORKBOOK_GENERATED',flush=True)
