from common import *
import sys,re,math,csv
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import load_workbook
import numpy as np

assert read(O/'results/cache_gate.json')['status']=='PASS_REAL_DEVELOPMENT_ALLOWED'
assert read(O/'results/development_gate.json')['status']=='STOP_DEVELOPMENT_GATE'
assert read(O/'results/performance_gate.json')['status']=='NOT_RUN_DEVELOPMENT_GATE_FAILED'
assert read(O/'reviews/model_and_cache_verification.json')['status']=='PASS'
protected=read(O/'reviews/protected_before.json')
for item in protected:assert identity(item['path'])==item,item['path']
frozen=0
for name in ['reviews/frozen_protocol_before_scale.json','reviews/frozen_before_training.json','models/frozen_after_training.json','reviews/frozen_before_deployment.json','reviews/frozen_configuration_repair.json','reviews/frozen_development_inputs.json','reviews/frozen_analysis_before_final_audit.json']:
    check_freeze(O/name);frozen+=len(read(O/name)['files'])
budget=read(O/'models/training_budget.json');process=read(O/'logs/training_process.json')
assert budget['status']==process['status']=='PASS' and budget['runs']==4 and budget['updates']==4800 and budget['all_runs_complete']
assert budget['elapsed_seconds']<600 and process['wall_seconds']<600
assert (O/'bin/evaluate').is_file() and (O/'logs/evaluation_jobs.json').is_file()
assert not (O/'reviews/frozen_validation_inputs.json').exists()
assert not list((O/'results').glob('validation_*'))
tables=read(O/'results/report_tables.json');wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);cells=0
for name,rows in tables.items():
    actual=list(wb[name].iter_rows(values_only=True));fields=list(rows[0]);assert list(actual[0])==fields
    assert len(actual)==len(rows)+1
    with (O/'results'/(name+'.csv')).open() as f:csvrows=list(csv.DictReader(f))
    assert len(csvrows)==len(rows)
    for record,values,csvrow in zip(rows,actual[1:],csvrows):
        for key,val in zip(fields,values):
            ref=record[key]
            if isinstance(ref,float):assert math.isclose(val,ref,rel_tol=1e-12,abs_tol=1e-12)
            else:assert val==ref,(name,key,val,ref)
            assert csvrow[key]==('' if ref is None else str(ref))
            cells+=1
sources=read(O/'reviews/final_plots_generated.json')['sources']
assert list(wb['Sources'].iter_rows(values_only=True))[1:]==[tuple(x) for x in sources]
cells+=len(sources)*2;wb.close()
for ds in DS:
    c=read(O/'inputs'/ds/'fit/config.json');n=c['num_queries'];rr=np.arange(n)
    hits=np.fromfile(E28/'labels'/ds/'hits.u8',np.uint8).reshape(n,56)
    ndc=np.fromfile(E28/'labels'/ds/'graph_ndc.u32',np.uint32).reshape(n,56)
    ts=np.fromfile(E28/'inputs'/ds/'proposals/triples.i32',np.int32).reshape(n,56,3)
    for method in ['anchor']+CELLS:
        z=np.load(O/'results'/ds/(method+'.npz'));slot=z['slot']
        assert np.array_equal(z['row_ids'],np.fromfile(c['row_ids'],np.int64))
        assert np.array_equal(z['selected_triples'],ts[rr,slot])
        assert np.array_equal(z['hits'],hits[rr,slot]) and np.array_equal(z['ndc'],ndc[rr,slot])
        m=next(x for x in tables['cache_metrics'] if x['dataset']==ds and x['method']==method and x['group']=='all')
        assert abs(z['hits'].mean()*10-m['greedy_recall_pct'])<1e-10
        assert abs(z['ndc'].mean()-m['greedy_graph_ndc'])<1e-9
    a,t,r=[next(x for x in tables['cache_metrics'] if x['dataset']==ds and x['method']==k and x['group']=='all') for k in ['anchor','teacher','reward']]
    conditions=[r['greedy_recall_pct']-a['greedy_recall_pct']>=.2-1e-10,r['greedy_recall_pct']-t['greedy_recall_pct']>=.05-1e-10,
        r['greedy_graph_ndc']<=a['greedy_graph_ndc']*1.01,r['expected_graph_ndc']<=a['expected_graph_ndc']*1.01]
    gate=next(x for x in read(O/'results/cache_gate.json')['by_dataset'] if x['dataset']==ds)
    assert list(gate['conditions'].values())==conditions and gate['passed']==all(conditions)
visual=read(O/'reviews/plot_visual_review.json');assert visual['status']=='PASS' and len(visual['figures'])==3
links=0
docs=[R/x for x in ['summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md','experiment_history/21_direct_reward_policy.md']]+[O/'README.md',O/'report.md',O/'protocol.md',O/'deployment_plan.md',O/'configuration_repair.md']
for doc in docs:
    for target in re.findall(r'\]\(([^)]+)\)',doc.read_text()):
        if '://' in target or target.startswith('#'):continue
        path=target.split('#')[0]
        assert (doc.parent/path).resolve().exists(),(doc,target)
        links+=1

assert read(O/'reviews/report_review.json')['status']=='PASS'
from deployment_common_v2 import verify_outputs
import deployment_common_v2 as deployment
def verify_existing_record(path,obj):
    assert read(path)==obj,path
deployment.dump=verify_existing_record
jobs=read(O/'logs/evaluation_jobs.json');actual=0;hnsw=read(O/'reviews/route_parity.json')['HNSW_queries'];query_outputs=0;point_count=0
for job in jobs:
    if job['status']=='FAILED':
        assert job['verified_actual_calls']==job['charged_calls']==0
        assert job['reserved_calls']==160064 and read(O/'reviews/configuration_repair.json')['status']=='PASS'
        continue
    assert job['status']=='COMPLETE' and job['exit_code']==0
    dest=O/'results'/(job['run_id']+'_'+job['dataset']);c=read(dest/'config.json');r=read(dest/'results.json')
    assert identity(job['config']['path'])==job['config']
    assert r['physical_graph_calls']==job['charged_calls']==job['reserved_calls']
    assert c['timing_repeats']==r['timing_repeats']==0
    verify_outputs(dest,c,r)
    actual+=r['physical_graph_calls'];query_outputs+=r['queries']*len(r['methods']);point_count+=len(r['methods'])
    hnsw+=sum(2*r['queries']*(2 if m['family']=='c8' else 1) for m in r['methods'])
    for m in r['methods']:assert m['qps'] is None and m['end_to_end_seconds']==[] and m['qps_runs']==[]
assert actual==309008 and hnsw==805056 and query_outputs==154336 and point_count==21
assert sum(x['reserved_calls'] for x in jobs)==469072<800000
assert sum(x.get('charged_calls',x['reserved_calls']) for x in jobs)==actual<800000
assert sum(x['seconds'] for x in jobs)<5400
for ds in DS:
    for split in ['cal','audit']:
        r=read(O/'results'/(split+'_'+ds)/'results.json');ms={m['model_name']:m for m in r['methods']}
        default='gt10_soft' if ds==DS[0] else 'gt10_hard';reward,teacher,base=[ms[k] for k in ['reward','teacher',default]]
        conditions=[reward['total_hits']>teacher['total_hits'],reward['total_hits']>base['total_hits'],reward['mean_graph_ndc']<=base['mean_graph_ndc']*1.01]
        item=next(x for x in read(O/'results/development_gate.json')['by_dataset_split'] if x['dataset']==ds and x['split']==split)
        assert list(item['conditions'].values())==conditions and item['passed']==all(conditions)

size=sum(p.stat().st_size for p in O.rglob('*') if p.is_file())
free=shutil.disk_usage(O).free;assert free>10*(1<<30)
result=dict(status='PASS',protected_identity_checks=len(protected),frozen_identity_checks=frozen,training_runs=4,optimizer_updates=4800,
    total_training_seconds=budget['elapsed_seconds'],new_labels=0,logical_labels_reused=2240000,real_search_calls=309008,new_HNSW_queries=805056,new_QPS_points=0,
    model_checkpoint_verifications=8,cache_models_scored=6,workbook_verified_cells=cells,figures_visually_reviewed=3,
    document_links_checked=links,artifact_bytes=size,free_bytes=free,heldout_access=False,timestamp=time.time())
dump(O/'reviews/final_audit.json',result)
dump(O/'state.json',dict(status='COMPLETE',conclusion='DIRECT_REWARD_POLICY_COMPLETE_FIT_GAIN_MIXED_DEVELOPMENT_STOP',training_runs=4,optimizer_updates=4800,
    real_search_calls=309008,new_labels=0,heldout_access=False,gate='STOP_DEVELOPMENT_GATE'))
print(json.dumps(result,indent=2),flush=True)
