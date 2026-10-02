from common import *
import sys
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

SHORT={DS[0]:'WebVid2.5M',DS[1]:'Yandex10M IP'}
COLORS={'teacher':'#327cb0','reward':'#d7762c'}
LABELS={'teacher':'Teacher hinge','reward':'Direct reward'}
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none'})
def save(fig,name):
    for ext in ['png','svg']:fig.savefig(O/'figures'/(name+'.'+ext),dpi=160,facecolor='white')
    plt.close(fig)

assert read(O/'state.json')['status']=='CACHE_ANALYSIS_COMPLETE'
t=read(O/'results/analysis_tables.json')
fig,axs=plt.subplots(3,2,figsize=(14.5,12))
for j,ds in enumerate(DS):
    base=next(x for x in t['cache_metrics'] if x['dataset']==ds and x['method']=='anchor' and x['group']=='all')
    for cell in CELLS:
        rows=[x for x in t['training_curve'] if x['dataset']==ds and x['cell']==cell]
        xs=[x['epoch'] for x in rows]
        model=next(x for x in t['model_comparison'] if x['dataset']==ds and x['cell']==cell)
        metric=next(x for x in t['cache_metrics'] if x['dataset']==ds and x['method']==cell and x['group']=='all')
        for i,key in enumerate(['expected_recall_pct','greedy_recall_pct','greedy_graph_ndc']):
            ys=[x[key]-rows[0][key] if i<2 else (x[key]/rows[0][key]-1)*100 for x in rows]
            axs[i,j].plot(xs,ys,color=COLORS[cell],label=LABELS[cell],lw=1.6)
            selected=metric[key]-base[key] if i<2 else (metric[key]/base[key]-1)*100
            axs[i,j].scatter([model['selected_epoch']],[selected],color=COLORS[cell],marker='D',s=38,edgecolor='#263544',lw=.6,zorder=5,
                label='Selected float32 export' if cell=='teacher' else None)
    for ax in axs[:,j]:ax.axhline(0,color='#72808a',lw=.7);ax.grid(alpha=.15);ax.set_xlabel('Optimizer updates')
    axs[0,j].set_title(SHORT[ds]+' — policy expectation')
    axs[1,j].set_title('Actual chosen triple: cached L128 labels')
    axs[2,j].set_title('Graph distance computations for chosen triple')
    axs[0,j].set_ylabel('Expected Recall change (pp)')
    axs[1,j].set_ylabel('Greedy Recall change (pp)')
    axs[2,j].set_ylabel('Graph NDC change (%)')
    axs[1,j].axhline(.2,color='#358e6e',ls=':',label='Anchor gain gate: +0.20 pp')
    axs[2,j].axhline(1,color='#9c6666',ls=':',label='Cost cap: +1%')
fig.suptitle('Can direct search reward improve the deployed choice?',fontsize=18,y=.98)
hs,ls=axs[1,0].get_legend_handles_labels();fig.legend(hs,ls,loc='lower center',bbox_to_anchor=(.5,.045),ncol=4,frameon=False,fontsize=10)
fig.text(.5,.02,'Fixed C8 / 56 complete triples; fit-only real-label cache. Lines: float64 trajectory. No new end-to-end QPS.',ha='center',fontsize=10,color='#586a76')
fig.subplots_adjust(left=.085,right=.98,top=.92,bottom=.12,hspace=.4,wspace=.25);save(fig,'reward_and_greedy')

fig,axs=plt.subplots(2,2,figsize=(14,8.8))
for j,ds in enumerate(DS):
    for cell in CELLS:
        rows=[x for x in t['training_curve'] if x['dataset']==ds and x['cell']==cell]
        axs[0,j].plot([x['epoch'] for x in rows],[(x['objective']/rows[0]['objective']-1)*100 for x in rows],color=COLORS[cell],label=LABELS[cell])
    rows=[next(x for x in t['paired_cache'] if x['dataset']==ds and x['method']==cell and x['reference']=='anchor' and x['group']=='all') for cell in CELLS]
    x=np.arange(2);width=.3
    axs[1,j].bar(x-width/2,[r['expected_delta_recall_pp'] for r in rows],width,color='#9cacb8',label='Expected Recall gain')
    ys=np.array([r['greedy_delta_recall_pp'] for r in rows]);lo=np.array([r['ci95_low_pp'] for r in rows]);hi=np.array([r['ci95_high_pp'] for r in rows])
    axs[1,j].bar(x+width/2,ys,width,color=[COLORS[c] for c in CELLS],label='Greedy Recall gain')
    axs[1,j].errorbar(x+width/2,ys,yerr=[ys-lo,hi-ys],fmt='none',color='#25394a',capsize=4)
    for k,v in enumerate(ys):axs[1,j].annotate(f'{v:+.4f}',(x[k]+width/2,v),xytext=(0,8 if v>=0 else -17),textcoords='offset points',ha='center',fontsize=10)
    axs[0,j].set_title(SHORT[ds]);axs[0,j].set_xlabel('Optimizer updates');axs[0,j].set_ylabel('Own total objective change (%)')
    axs[1,j].set_xticks(x,[LABELS[c] for c in CELLS]);axs[1,j].set_ylabel('Recall gain vs C8 anchor (pp)')
    axs[1,j].axhline(.2,color='#358e6e',ls=':',label='Anchor gain gate: +0.20 pp')
    axs[1,j].axhline(0,color='#72808a',lw=.7)
    for ax in axs[:,j]:ax.grid(axis='y',alpha=.15)
    axs[0,j].legend(frameon=False,fontsize=10)
fig.suptitle('Optimization and final cached utility are separate checks',fontsize=18,y=.98)
fig.legend(*axs[1,0].get_legend_handles_labels(),loc='lower center',bbox_to_anchor=(.5,.065),ncol=3,frameon=False,fontsize=10)
fig.text(.5,.037,'Each objective is normalized to its own start; absolute values across objectives are not comparable.',ha='center',fontsize=9,color='#586a76')
fig.text(.5,.012,'Whiskers: 2,000 paired fit-query bootstrap draws, descriptive only. Additional gate: +0.05 pp over teacher on both datasets.',ha='center',fontsize=9,color='#586a76')
fig.subplots_adjust(left=.08,right=.98,top=.91,bottom=.18,hspace=.34,wspace=.25);save(fig,'objective_and_final_utility')

wb=Workbook();wb.remove(wb.active)
for name,rows in t.items():
    ws=wb.create_sheet(name);fields=list(rows[0]);ws.append(fields)
    for row in rows:ws.append([row.get(k) for k in fields])
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='263C50')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(42,max(14,len(str(col[0].value))+2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value,float):c.number_format='0.000000'
sources=[('Protocol',str(O/'protocol.md')),('Raw complete-action labels',str(E28/'labels')),('Frozen scales',str(O/'reviews/reward_scales.json')),
    ('Report',str(O/'report.md')),('Expected reward','Exact softmax expectation over all 56 labeled C8 triples; not greedy deployment'),
    ('Greedy utility','Selected float32 model argmax triple; cached real NSG L128/K10 labels'),
    ('Scope','Historical fit20K, one H/seed/lambda/temperature; no new QPS or independent blind inference'),
    ('Cost','Graph NDC feasibility only; any future online timer must include complete C8 candidate generation')]
ws=wb.create_sheet('Sources');ws.append(['Item','Definition / file'])
for row in sources:ws.append(row)
ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=125
wb.save(O/'chart_data.xlsx')
dump(O/'reviews/plots_generated.json',dict(status='PASS',figures=2,formats=['png','svg'],workbook_sheets=5,sources=sources))
print('PLOTS_AND_XLSX_COMPLETE',flush=True)
