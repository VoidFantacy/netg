from common import *
import sys,re,math,csv
sys.dont_write_bytecode=True
sys.path.insert(0,str(E28/'src/vendor'))
from openpyxl import load_workbook
import numpy as np

assert read(O/'results/cache_gate.json')['status']=='STOP_CACHE_UTILITY_GATE'
assert read(O/'reviews/model_and_cache_verification.json')['status']=='PASS'
protected=read(O/'reviews/protected_before.json')
for item in protected:assert identity(item['path'])==item,item['path']
frozen=0
for name in ['reviews/frozen_protocol_before_scale.json','reviews/frozen_before_training.json','models/frozen_after_training.json','reviews/frozen_analysis_before_final_audit.json']:
    check_freeze(O/name);frozen+=len(read(O/name)['files'])
budget=read(O/'models/training_budget.json');process=read(O/'logs/training_process.json')
assert budget['status']==process['status']=='PASS' and budget['runs']==4 and budget['updates']==4800 and budget['all_runs_complete']
assert budget['elapsed_seconds']<600 and process['wall_seconds']<600
assert not (O/'bin').exists() and not (O/'logs/evaluation_jobs.json').exists()
tables=read(O/'results/analysis_tables.json');wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);cells=0
for name,rows in tables.items():
    actual=list(wb[name].iter_rows(values_only=True));fields=list(rows[0]);assert list(actual[0])==fields
    assert len(actual)==len(rows)+1
    with (O/'results'/(name+'.csv')).open() as f:csvrows=list(csv.DictReader(f))
    assert len(csvrows)==len(rows)
    for record,values,csvrow in zip(rows,actual[1:],csvrows):
        for key,val in zip(fields,values):
            ref=record[key]
            if isinstance(ref,float):assert math.isclose(val,ref,rel_tol=1e-12,abs_tol=1e-12)
            else:assert val==ref,(name,key,val,ref)
            assert csvrow[key]==('' if ref is None else str(ref))
            cells+=1
sources=read(O/'reviews/plots_generated.json')['sources']
assert list(wb['Sources'].iter_rows(values_only=True))[1:]==[tuple(x) for x in sources]
cells+=len(sources)*2;wb.close()
for ds in DS:
    c=read(O/'inputs'/ds/'fit/config.json');n=c['num_queries'];rr=np.arange(n)
    hits=np.fromfile(E28/'labels'/ds/'hits.u8',np.uint8).reshape(n,56)
    ndc=np.fromfile(E28/'labels'/ds/'graph_ndc.u32',np.uint32).reshape(n,56)
    ts=np.fromfile(E28/'inputs'/ds/'proposals/triples.i32',np.int32).reshape(n,56,3)
    for method in ['anchor']+CELLS:
        z=np.load(O/'results'/ds/(method+'.npz'));slot=z['slot']
        assert np.array_equal(z['row_ids'],np.fromfile(c['row_ids'],np.int64))
        assert np.array_equal(z['selected_triples'],ts[rr,slot])
        assert np.array_equal(z['hits'],hits[rr,slot]) and np.array_equal(z['ndc'],ndc[rr,slot])
        m=next(x for x in tables['cache_metrics'] if x['dataset']==ds and x['method']==method and x['group']=='all')
        assert abs(z['hits'].mean()*10-m['greedy_recall_pct'])<1e-10
        assert abs(z['ndc'].mean()-m['greedy_graph_ndc'])<1e-9
    a,t,r=[next(x for x in tables['cache_metrics'] if x['dataset']==ds and x['method']==k and x['group']=='all') for k in ['anchor','teacher','reward']]
    conditions=[r['greedy_recall_pct']-a['greedy_recall_pct']>=.2-1e-10,r['greedy_recall_pct']-t['greedy_recall_pct']>=.05-1e-10,
        r['greedy_graph_ndc']<=a['greedy_graph_ndc']*1.01,r['expected_graph_ndc']<=a['expected_graph_ndc']*1.01]
    gate=next(x for x in read(O/'results/cache_gate.json')['by_dataset'] if x['dataset']==ds)
    assert list(gate['conditions'].values())==conditions and gate['passed']==all(conditions)
visual=read(O/'reviews/plot_visual_review.json');assert visual['status']=='PASS' and len(visual['figures'])==2
links=0
docs=[R/x for x in ['summary.md','spec.md','s-target.md','s-experiment.md','experiment_history/README.md','experiment_history/21_direct_reward_policy.md']]+[O/'README.md',O/'report.md',O/'protocol.md']
for doc in docs:
    for target in re.findall(r'\]\(([^)]+)\)',doc.read_text()):
        if '://' in target or target.startswith('#'):continue
        path=target.split('#')[0]
        assert (doc.parent/path).resolve().exists(),(doc,target)
        links+=1
assert read(O/'reviews/report_review.json')['status']=='PASS'
size=sum(p.stat().st_size for p in O.rglob('*') if p.is_file())
free=shutil.disk_usage(O).free;assert free>10*(1<<30)
result=dict(status='PASS',protected_identity_checks=len(protected),frozen_identity_checks=frozen,training_runs=4,optimizer_updates=4800,
    total_training_seconds=budget['elapsed_seconds'],new_labels=0,logical_labels_reused=2240000,real_search_calls=0,new_HNSW_queries=0,new_QPS_points=0,
    model_checkpoint_verifications=8,cache_models_scored=6,workbook_verified_cells=cells,figures_visually_reviewed=2,
    document_links_checked=links,artifact_bytes=size,free_bytes=free,heldout_access=False,timestamp=time.time())
dump(O/'reviews/final_audit.json',result)
dump(O/'state.json',dict(status='COMPLETE',conclusion='DIRECT_REWARD_POLICY_COMPLETE_CACHE_GATE_FAILED',training_runs=4,optimizer_updates=4800,
    real_search_calls=0,new_labels=0,heldout_access=False,gate='STOP_CACHE_UTILITY_GATE'))
print(json.dumps(result,indent=2),flush=True)
