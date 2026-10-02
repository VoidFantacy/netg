from common import *
import csv,re,sys,datetime
sys.path.insert(0,str(R/'entry_method_exploration/20260911_gt10_shared_transform_v1/src/vendor'))
from openpyxl import load_workbook

assert not (O/'reviews/final_audit.json').exists()
assert read(O/'reviews/visual_review.json')['status']=='PASS'
assert read(O/'reviews/output_verification.json')['status']=='PASS'
jobs=read(O/'logs/jobs.json');assert len(jobs)==2 and all(j['status']=='COMPLETE' and j['exit_code']==0 and j['process_threads_max']==1 for j in jobs)
assert sum(j['actual_calls'] for j in jobs)==630720==sum(j['reserved_calls'] for j in jobs)
for name in ['upstream.json','E33_preserved.json','before_queries.json','regression_sources.json','analysis_implementation.json']:check(O/'reviews'/name)
a=read(O/'results/analysis.json');assert len(a['points'])==45 and len(a['qps_runs'])==225 and len(a['paired'])==35
mapping={'Points':'points','Timing_runs':'qps_runs','Paired_changes':'paired','Measured_frontier':'frontiers','Same_recall_estimates':'same_recall_interpolation'}
wb=load_workbook(O/'chart_data.xlsx',read_only=True,data_only=True);cells=0
for sheet,key in mapping.items():
 rows=list(wb[sheet].values);expected=a[key];keys=list(expected[0]);assert list(rows[0])==keys and len(rows)==len(expected)+1
 for got,row in zip(rows[1:],expected):
  for v,k in zip(got,keys):
   e=row[k]
   if isinstance(e,(int,float)) and not isinstance(e,bool):assert v is not None and abs(v-e)<1e-9*max(1,abs(e)),(sheet,k,v,e)
   else:assert v==e,(sheet,k,v,e)
   cells+=1
csvmap={'points.csv':'points','qps_runs.csv':'qps_runs','paired_changes.csv':'paired','frontier_points.csv':'frontiers','same_recall_interpolation.csv':'same_recall_interpolation'}
for file,key in csvmap.items():
 with (O/'results'/file).open() as f:rows=list(csv.DictReader(f))
 assert len(rows)==len(a[key])
 for got,row in zip(rows,a[key]):
  assert got=={k:'' if v is None else str(v) for k,v in row.items()}
for figure in ['recall_qps','metrics_by_search_L','rl_changes']:
 for ext in ['png','svg','pdf']:assert (O/'figures'/f'{figure}.{ext}').stat().st_size>1000
docs=[R/x for x in ['AGENTS.md','summary.md','spec.md','s-target.md','s-experiment.md']]+[R/'experiment_history/README.md',R/'experiment_history/23_hnsw_rl_benchmark.md',O/'README.md',O/'protocol.md',O/'report.md']
links=0
for p in docs:
 for link in re.findall(r'\]\(([^)]+)\)',p.read_text()):
  link=link.strip('<>').split('#')[0]
  if not link or '://' in link:continue
  assert (p.parent/link).exists(),(p,link)
  links+=1
assert (R/'AGENTS.md').read_bytes()==(O/'source_before/AGENTS.md').read_bytes()
resources()
freeze(O/'reviews/frozen_completion.json',docs+[p for p in (O/'src').glob('*') if p.is_file()]+[p for p in (O/'figures').iterdir() if p.is_file()]+[O/'chart_data.xlsx',O/'results/analysis.json']+list((O/'results').glob('*.csv')))
model=[x.split(':',1)[1].strip() for x in Path('/proc/cpuinfo').read_text().splitlines() if x.startswith('model name')][0]
out=dict(status='PASS',points=45,wall_measurements=225,all_GT_rescored_query_outputs=90000,old_baseline_bit_regression_outputs=70000,
 actual_NSG_calls=630720,reserved_calls=630720,repair_calls=0,optimizer_updates=0,E33_preserved_files=len(read(O/'reviews/E33_preserved.json')['files']),
 upstream_SHA_files=len(read(O/'reviews/upstream.json')['files']),all_frozen_sources_and_inputs_unchanged=True,workbook_verified_cells=cells,workbook_sheets=wb.sheetnames,
 visually_reviewed_figures=3,local_links_checked=links,cpu_model=model,cpu=15,process_OMP_BLAS_MKL_threads=1,validation_rows_per_dataset=2000,validation_historical_use=True,heldout_read_or_hash=False,
 artifact_bytes=sum(p.stat().st_size for p in O.rglob('*') if p.is_file()),free_disk_bytes=shutil.disk_usage(O).free)
dump(O/'reviews/final_audit.json',out)
dump(O/'state.json',dict(status='COMPLETE_FROZEN_RL_BENCHMARK',final_audit_status='PASS',points=45,wall_measurements=225,actual_NSG_calls=630720,
 optimizer_updates=0,model_reselection=False,default_replaced=False,heldout_access=False,closed_at=datetime.datetime.now(datetime.timezone.utc).isoformat()))
print('FINAL_AUDIT_PASS',json.dumps(out),flush=True)
