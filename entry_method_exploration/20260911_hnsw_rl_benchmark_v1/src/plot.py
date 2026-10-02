from common import *
import sys
sys.path.insert(0,str(R/'entry_method_exploration/20260911_gt10_shared_transform_v1/src/vendor'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from openpyxl import Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter

assert read(O/'reviews/output_verification.json')['status']=='PASS'
a=read(O/'results/analysis.json')
COLORS={'nsg':'#777777','fast30k':'#3268b2','gt10_hard':'#ce8a13','gt10_soft':'#259281','rl':'#c23672'}
MARKERS={'nsg':'s','fast30k':'^','gt10_hard':'D','gt10_soft':'v','rl':'o'}
TITLE={DS[0]:'WebVid-2.5M',DS[1]:'Yandex-10M / IP'}
ORDER=list(LABELS)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','svg.fonttype':'none','savefig.facecolor':'white'})
handles=[Line2D([0],[0],color=COLORS[m],marker=MARKERS[m],lw=2,label=LABELS[m]) for m in ORDER]

def data(ds,m):return sorted([p for p in a['points'] if p['dataset']==ds and p['method']==m],key=lambda p:p['L'])
def style(ax):ax.grid(alpha=.18);ax.set_axisbelow(True)
def save(fig,name):
 for ext in ['svg','png','pdf']:
  path=O/'figures'/f'{name}.{ext}';assert not path.exists();fig.savefig(path,dpi=180,bbox_inches='tight')
 plt.close(fig)

fig,axs=plt.subplots(1,2,figsize=(15,5.8))
for ax,ds in zip(axs,DS):
 for m in ORDER:
  ps=data(ds,m)
  if not ps:continue
  x=[p['recall_pct'] for p in ps];y=[p['qps'] for p in ps]
  ax.plot(x,y,color=COLORS[m],marker=MARKERS[m],lw=2.5 if m=='rl' else 1.6,ms=6,zorder=5 if m=='rl' else 3)
  ax.fill_between(x,[p['qps_min'] for p in ps],[p['qps_max'] for p in ps],color=COLORS[m],alpha=.12)
  if m=='rl':
   for p in ps:ax.annotate(f"L={p['L']}",(p['recall_pct'],p['qps']),xytext=(5,6),textcoords='offset points',fontsize=9,color=COLORS[m])
 ax.set(title=TITLE[ds],xlabel='Recall@10 (%)',ylabel='Queries / second');ax.set_yscale('log');style(ax)
fig.suptitle('Frozen HNSW-RL versus NSG, Fast and FastGT10',fontsize=17,y=1.04)
fig.legend(handles=handles,loc='upper center',bbox_to_anchor=(.5,.985),ncol=5,frameon=False)
fig.text(.5,-.035,'Historical validation: 2,000 queries / dataset; CPU 15, one thread. Lines: measured points; bands: min-max of five QPS runs.',ha='center',fontsize=10)
fig.subplots_adjust(top=.82,wspace=.23);save(fig,'recall_qps')

fig,axs=plt.subplots(2,3,figsize=(16,9))
fields=[('qps','QPS (higher is better)'),('recall_pct','Recall@10 % (higher is better)'),('failure_pct','Zero-hit queries % (lower is better)')]
for row,ds in enumerate(DS):
 for col,(field,label) in enumerate(fields):
  ax=axs[row,col]
  for m in ORDER:
   ps=data(ds,m)
   if not ps:continue
   xs=[p['L'] for p in ps];ys=[p[field] for p in ps]
   ax.plot(xs,ys,color=COLORS[m],marker=MARKERS[m],lw=2.6 if m=='rl' else 1.7,ms=5,zorder=5 if m=='rl' else 3)
   if field=='qps':ax.fill_between(xs,[p['qps_min'] for p in ps],[p['qps_max'] for p in ps],color=COLORS[m],alpha=.12)
  ax.set_xscale('log',base=2);ax.set_xticks(LS,labels=LS);ax.set(xlabel='NSG search L',ylabel=label,title=TITLE[ds]);style(ax)
  if field=='qps':ax.set_yscale('log')
  if field=='failure_pct':ax.set_ylim(bottom=0)
fig.suptitle('QPS, recall and query failure rate across search budgets',fontsize=17,y=1.02)
fig.legend(handles=handles,loc='upper center',bbox_to_anchor=(.5,.982),ncol=5,frameon=False)
fig.text(.5,-.025,'Failure = Recall@10 < 0.1 = zero exact top-10 hits. QPS uses five independent batch walls; shaded range is timing variation.',ha='center',fontsize=10)
fig.subplots_adjust(top=.87,hspace=.38,wspace=.28);save(fig,'metrics_by_search_L')

fig,axs=plt.subplots(2,3,figsize=(16,9))
fields=[('qps_change_pct','RL QPS change (%)',None,None),('recall_delta_pp','RL recall change (percentage points)','recall_ci_low_pp','recall_ci_high_pp'),('failure_delta_pp','RL failure change (percentage points)','failure_ci_low_pp','failure_ci_high_pp')]
for row,ds in enumerate(DS):
 for col,(field,label,lo,hi) in enumerate(fields):
  ax=axs[row,col]
  for m in ORDER:
   if m=='rl':continue
   ps=sorted([p for p in a['paired'] if p['dataset']==ds and p['baseline']==m],key=lambda p:p['L'])
   if not ps:continue
   x=np.array([p['L'] for p in ps]);y=np.array([p[field] for p in ps]);ax.plot(x,y,color=COLORS[m],marker=MARKERS[m],lw=1.7,ms=5)
   if lo:ax.fill_between(x,[p[lo] for p in ps],[p[hi] for p in ps],color=COLORS[m],alpha=.07)
  ax.axhline(0,color='#333333',lw=.8,ls='--');ax.set_xscale('log',base=2);ax.set_xticks(LS,labels=LS)
  ax.set(title=TITLE[ds],xlabel='NSG search L',ylabel=label);style(ax)
fig.suptitle('HNSW-RL minus each baseline at the same search L',fontsize=17,y=1.02)
fig.legend(handles=handles[:-1],loc='upper center',bbox_to_anchor=(.5,.982),ncol=4,frameon=False)
fig.text(.5,-.025,'Positive: faster QPS / higher recall. Negative failure change: fewer zero-hit queries. Recall/failure bands: paired query-bootstrap 95% intervals.',ha='center',fontsize=10)
fig.subplots_adjust(top=.87,hspace=.38,wspace=.28);save(fig,'rl_changes')

wb=Workbook();wb.remove(wb.active)
tables={'Points':a['points'],'Timing_runs':a['qps_runs'],'Paired_changes':a['paired'],'Measured_frontier':a['frontiers'],'Same_recall_estimates':a['same_recall_interpolation']}
tables['Methods']=[dict(method=m,label=LABELS[m],note=('original Fast30K, graph seed100 and legacy heap order' if m=='fast30k' else 'WebVid soft-t1 / Yandex hard starting graph; E33 selected physical graph, no retraining' if m=='rl' else 'WebVid RL direct starting graph' if m=='gt10_soft' else 'fixed original NSG entry' if m=='nsg' else 'GT10 hard, fit20K')) for m in ORDER]
tables['Sources']=[dict(artifact=str(O/'protocol.md'),meaning='frozen measurement protocol'),dict(artifact=str(O/'results/analysis.json'),meaning='authoritative numerical table data'),dict(artifact=str(O/'reviews/output_verification.json'),meaning='independent output and regression checks'),dict(artifact=str(E33/'report.md'),meaning='prior RL training and selected model')]
for idx,(name,rows) in enumerate(tables.items()):
 ws=wb.create_sheet(name);keys=list(rows[0]);ws.append(keys)
 for row in rows:ws.append([row[k] for k in keys])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=__import__('openpyxl').styles.Font(bold=True,color='FFFFFF');c.fill=__import__('openpyxl').styles.PatternFill('solid',fgColor='244B67')
 for col in range(1,len(keys)+1):ws.column_dimensions[get_column_letter(col)].width=min(55,max(15,len(keys[col-1])+2))
 tab=Table(displayName=f'E34Table{idx+1}',ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
assert not (O/'chart_data.xlsx').exists();wb.save(O/'chart_data.xlsx')
dump(O/'reviews/plot_sources.json',dict(status='GENERATED_PENDING_VISUAL_REVIEW',figures=['recall_qps','metrics_by_search_L','rl_changes'],source=str(O/'results/analysis.json'),xlsx_sheets=list(tables)))
print('PLOTS_AND_WORKBOOK_SAVED',flush=True)
