from setup import *
import numpy as np,csv,sys,itertools,collections
sys.path.insert(0,str(O/'src/vendor'))
from openpyxl import Workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
NAMES={'nsg':'NSG','fast30k':'Fast30K','gt1_hard':'GT1 hard','gt10_hard':'GT10 hard','gt1_soft':'GT1 soft','gt10_soft':'GT10 soft','linear':'Shared linear','mlp':'Small MLP','free':'Free vectors'}
SHORT={DS[0]:'WebVid2.5M',DS[1]:'Yandex10M IP'}
COL={'nsg':'#9ca3af','fast30k':'#8a8b2a','gt1_hard':'#c084fc','gt10_hard':'#243b53','gt1_soft':'#bbb0bf','gt10_soft':'#9a6d47','linear':'#118a8c','mlp':'#e76f51','free':'#497cc5'}
def writecsv(name,rows):
 path=O/'results'/f'{name}.csv'
 if not rows:return
 fields=list(dict.fromkeys(k for row in rows for k in row))
 with path.open('w') as f:
  w=csv.DictWriter(f,fields);w.writeheader();w.writerows(rows)
def values(dest,m,name,dtype):return np.fromfile(dest/m['name']/name,dtype)
def bootstrap_delta(diff):
 rng=np.random.default_rng(20260911);means=[]
 for _ in range(20):means.extend(diff[rng.integers(0,len(diff),(100,len(diff)))].mean(1)*10)
 return list(map(float,np.quantile(means,[.025,.975])))
def frontier(ms):
 return sorted([m for m in ms if not any(n['recall_at_10']>=m['recall_at_10'] and n['qps']>=m['qps'] and (n['recall_at_10']>m['recall_at_10'] or n['qps']>m['qps']) for n in ms)],key=lambda m:m['recall_at_10'])
def interp(ms,r,rep=None):
 ms=frontier(ms);xs=np.array([m['recall_at_10'] for m in ms]);ys=np.array([m['qps'] if rep is None else m['qps_runs'][rep] for m in ms])
 if r<xs.min() or r>xs.max():return None
 return float(np.exp(np.interp(r,xs,np.log(ys))))
def main():
 assert json.loads((O/'state.json').read_text())['status']=='EVALUATION_COMPLETE_ANALYSIS_PENDING'
 allrows=[];checks=[];comparisons=[];interpolated=[];diagnostics=[];val={}
 for dest in sorted((O/'results').iterdir()):
  if not dest.is_dir() or not (dest/'results.json').exists():continue
  res=json.loads((dest/'results.json').read_text());c=json.loads((dest/'config.json').read_text());n=c['num_queries'];ds=c['dataset'];stage=dest.name.removesuffix('_'+ds)
  gt=np.fromfile(c['gt_file'],np.int64).reshape(n,100)[:,:10]
  for m in res['methods']:
   ids=values(dest,m,'ids.i32',np.int32).reshape(n,10);hits=values(dest,m,'hits.u8',np.uint8);masks=values(dest,m,'hitmask.u16',np.uint16)
   member=(ids[:,:,None]==gt[:,None,:]).any(1);rh=member.sum(1);rm=(member*(1<<np.arange(10,dtype=np.uint16))).sum(1)
   assert np.array_equal(rh,hits) and np.array_equal(rm,masks);assert hits.sum()==m['total_hits'] and int(hits.sum())/(n*10)==m['recall_at_10']
   assert (ids>=0).all() and (np.sort(ids,axis=1)[:,1:]!=np.sort(ids,axis=1)[:,:-1]).all()
   entries=values(dest,m,'entry_ids.i32',np.int32).reshape(n,3)
   if m['configuration']['k']==3:assert (np.sort(entries,axis=1)[:,1:]!=np.sort(entries,axis=1)[:,:-1]).all()
   for label in ['route','graph']:
    a=values(dest,m,f'{label}_ndc.u32',np.uint32);assert a.mean()==m[f'mean_{label}_ndc']
   assert m['full_plain_count_parity_queries']==n and m['original_kernel_parity_queries']==16 and m['repeated_output_parity']
   if res['timing_repeats']:assert len(m['qps_runs'])==3 and np.median(m['qps_runs'])==m['qps']
   row=dict(dataset=SHORT[ds],stage=stage,method=m['model_name'],L=m['L'],exact=m['configuration'].get('exact',False),queries=n,recall_pct=m['recall_at_10']*100,total_hits=m['total_hits'],qps=m['qps'],qps_min=m['qps_min'],qps_max=m['qps_max'],route_ndc=m['mean_route_ndc'],graph_ndc=m['mean_graph_ndc'],total_ndc=m['mean_total_ndc'],failed=m['failed_queries'],build_seconds=m['policy_build_seconds'],route_us=None if m['route_only_seconds'] is None else m['route_only_seconds']/n*1e6,graph_us=None if m['graph_only_seconds'] is None else m['graph_only_seconds']/n*1e6,physical_calls=m['physical_graph_calls'])
   for i,q in enumerate(m['qps_runs']):row[f'qps_repeat_{i+1}']=q;row[f'wall_seconds_repeat_{i+1}']=m['end_to_end_seconds'][i]
   allrows.append(row)
  checks.append(dict(stage=stage,dataset=ds,points=len(res['methods']),queries=sum(m['queries'] for m in res['methods']),status='PASS',physical_calls=res['physical_graph_calls']))
  if stage=='validation':
   val[ds]=res['methods'];models={s:[m for m in res['methods'] if m['model_name']==s] for s in NAMES}
   for name,ms in models.items():
    for m in ms:
     if name=='nsg':continue
     entries=values(dest,m,'entry_ids.i32',np.int32);other=next(x for x in ms if x['L']==64);assert np.array_equal(entries,values(dest,other,'entry_ids.i32',np.int32))
   # Four original profiles plus NSG/Fast have deterministic historical outputs.
   e24res=json.loads((E24/'results'/ds/'results.json').read_text())
   oldname={'nsg':'nsg_medoid','fast30k':'fast30k_3','gt1_hard':'gt1_hard','gt10_hard':'gt10_hard','gt1_soft':'gt1_soft','gt10_soft':'gt10_soft'}
   for s,f in oldname.items():
    for m in models[s]:
     old=next(x for x in e24res['methods'] if x['family']==f and x['L']==m['L'])
     for filename in ['ids.i32','entry_ids.i32','graph_ndc.u32','hits.u8','hitmask.u16']:
      assert (dest/m['name']/filename).read_bytes()==(E24/'results'/ds/old['name']/filename).read_bytes(),(ds,s,m['L'],filename)
   for s in ['linear','mlp','free','gt1_hard','gt1_soft','gt10_soft']:
    for m in models[s]:
     b=next(x for x in models['gt10_hard'] if x['L']==m['L']);bh=values(dest,b,'hits.u8',np.uint8).astype(int);mh=values(dest,m,'hits.u8',np.uint8).astype(int);diff=mh-bh;ci=bootstrap_delta(diff)
     eq=(values(dest,b,'entry_ids.i32',np.int32).reshape(n,3)==values(dest,m,'entry_ids.i32',np.int32).reshape(n,3)).all(1)
     assert np.array_equal(values(dest,b,'ids.i32',np.int32).reshape(n,10)[eq],values(dest,m,'ids.i32',np.int32).reshape(n,10)[eq])
     comparisons.append(dict(dataset=SHORT[ds],method=s,L=m['L'],delta_recall_pp=float(diff.mean()*10),ci_low_pp=ci[0],ci_high_pp=ci[1],delta_qps_pct=(m['qps']/b['qps']-1)*100,delta_graph_ndc_pct=(m['mean_graph_ndc']/b['mean_graph_ndc']-1)*100,improved_queries=int((diff>0).sum()),harmed_queries=int((diff<0).sum()),gained_hits=int(diff[diff>0].sum()),lost_hits=int(-diff[diff<0].sum()),entry_change_pct=float((~eq).mean()*100),rescued_zero=int(((bh==0)&(mh>0)).sum()),new_zero=int(((bh>0)&(mh==0)).sum()),same_L_both_better=m['qps']>b['qps'] and m['recall_at_10']>b['recall_at_10']))
   # Original anchor recalls and predeclared gate targets, without extrapolation.
   targets=sorted(set([.90,.95,.98]+[m['recall_at_10'] for m in models['gt10_hard']]))
   for s in NAMES:
    if s=='gt10_hard':continue
    for r in targets:
     b=interp(models['gt10_hard'],r);q=interp(models[s],r)
     if b is None or q is None:continue
     ratios=[interp(models[s],r,k)/interp(models['gt10_hard'],r,k) for k in range(3)]
     interpolated.append(dict(dataset=SHORT[ds],method=s,recall_target_pct=r*100,baseline_qps=b,method_qps=q,delta_qps_pct=(q/b-1)*100,ratio_min=min(ratios),ratio_max=max(ratios),predeclared_gate_target=r in [.9,.95,.98],is_interpolation=True))
  # Training/cal/audit L128 means; audit exact-vs-HNSW is descriptive only.
  if stage in ['cal','audit','fit_probe']:
   for m in res['methods']:
    if m['L']==128:diagnostics.append(dict(dataset=SHORT[ds],stage=stage,method=m['model_name'],exact=m['configuration'].get('exact',False),recall_pct=m['recall_at_10']*100,graph_ndc=m['mean_graph_ndc']))
 for name,rows in [('all_points',allrows),('validation_differences',comparisons),('interpolated_same_recall',interpolated),('development_diagnostics',diagnostics)]:writecsv(name,rows)
 gate=[]
 for name in ['linear','mlp','free']:
  dsrecords=[]
  for ds in DS:
   xs=[r for r in interpolated if r['dataset']==SHORT[ds] and r['method']==name and r['predeclared_gate_target']]
   passing=[r for r in xs if r['delta_qps_pct']>=3 and r['ratio_min']>1]
   regressions=[r for r in interpolated if r['dataset']==SHORT[ds] and r['method']==name and r['delta_qps_pct']<=-3 and r['ratio_max']<1]
   dsrecords.append(dict(dataset=SHORT[ds],covered_targets=len(xs),passing_targets=len(passing),pass_at_least_two=len(passing)>=2,material_frontier_regressions=len(regressions)))
  applicable=all(x['covered_targets']>=2 for x in dsrecords)
  gate.append(dict(method=name,datasets=dsrecords,gate_applicable=applicable,gate_interpretation='APPLICABLE' if applicable else 'INSUFFICIENT_TARGET_COVERAGE_NOT_A_METHOD_FALSIFICATION',dual_dataset_GO=all(x['pass_at_least_two'] and not x['material_frontier_regressions'] for x in dsrecords)))
 dump(O/'reviews/output_verification.json',dict(status='PASS',jobs=checks,total_points=sum(x['points'] for x in checks),total_query_outputs=sum(x['queries'] for x in checks),historical_regression_outputs=120000,GT_rescore=True,same_entries_same_output=True,across_L_entries_equal=True))
 dump(O/'results/performance_gate.json',dict(scope='Predeclared coarse-grid interpolated necessary condition; not fresh validation or measured matched-recall certificate',methods=gate))
 make_plots(val,allrows,comparisons,interpolated)
 book=Workbook();book.remove(book.active)
 tables={'all_points':allrows,'delta_vs_GT10':comparisons,'interpolation':interpolated,'development':diagnostics}
 td=json.loads((O/'results/teacher_diagnostics.json').read_text());tables['teacher_groups']=td['groups'];tables['fit_known_subset']=td['model_labeled_subset']
 training=json.loads((O/'models/training_runs.json').read_text());tr=[]
 for m in training:tr.append(dict(dataset=SHORT[m['dataset']],family=m['family'],lambda_anchor=m['lambda_anchor'],parameters=m['parameter_count'],seconds=m['training_seconds'],export_seconds=m['export_seconds'],initial_CE=m['first']['ce'],final_CE=m['last']['ce'],final_objective=m['last']['objective'],final_grad_norm=m['last']['gradient_norm'],drift=m['last']['drift'],mean_cosine=m['cosine_to_anchor']['mean']))
 tables['training']=tr;writecsv('training_summary',tr)
 curves=[]
 for m in training:
  for point in json.loads((O/'models'/m['dataset']/f"{m['family']}_r{m['lambda_anchor']}"/'training_curve.json').read_text()):curves.append(dict(dataset=SHORT[m['dataset']],family=m['family'],lambda_anchor=m['lambda_anchor'],**point))
 tables['training_curves']=curves;writecsv('training_curve_data',curves)
 for name,rows in tables.items():
  sheet=book.create_sheet(name)
  if not rows:continue
  fields=list(dict.fromkeys(k for r in rows for k in r));sheet.append(fields)
  for r in rows:sheet.append([r.get(k) for k in fields])
  sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(32,max(14,len(str(col[0].value))+2))
 sheet=book.create_sheet('Sources');sheet.append(['Source','Location','Scope'])
 sheet.append(['Frozen experiment protocol','protocol.md','Models, selection, calls, timing and limits'])
 sheet.append(['Raw measured outputs','results/validation_<dataset>/results.json','QPS: original raw-query end-to-end; 3 repeats'])
 sheet.append(['Raw hit arrays','results/<stage>_<dataset>/<method>/hits.u8','Paired query bootstrap uses the same query rows'])
 sheet.append(['Training curves','models/<dataset>/<family>_r<lambda>/training_curve.json','All 81 checkpoints including epoch0; 80 updates'])
 sheet.append(['C8 teachers','labels/<dataset>/teacher.json','Fit-only, max hits then min graph NDC'])
 for x in json.loads((O/'sources/inventory.json').read_text()):sheet.append([x['name'],x['url'],'Primary paper; methodology only, no imported performance'])
 sheet.column_dimensions['A'].width=28;sheet.column_dimensions['B'].width=90;sheet.column_dimensions['C'].width=75;sheet.freeze_panes='A2'
 book.save(O/'chart_data.xlsx');print(json.dumps(dict(status='ANALYSIS_COMPLETE',points=len(allrows),gates=gate)))
def savefig(fig,name):
 fig.savefig(O/'figures'/f'{name}.png',dpi=190,bbox_inches='tight');fig.savefig(O/'figures'/f'{name}.svg',bbox_inches='tight');plt.close(fig)
def make_plots(val,rows,comps,interpolated):
 plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.grid':True,'grid.alpha':.2,'axes.titleweight':'bold'})
 for fname,names in [('recall_qps_main',['gt10_hard','linear','mlp','free']),('recall_qps_all',list(NAMES))]:
  fig,axs=plt.subplots(1,2,figsize=(13,5.6))
  for ax,ds in zip(axs,DS):
   for name in names:
    ms=sorted([m for m in val[ds] if m['model_name']==name],key=lambda m:m['L']);x=[100*m['recall_at_10'] for m in ms];y=[m['qps'] for m in ms]
    ax.errorbar(x,y,yerr=[[m['qps']-m['qps_min'] for m in ms],[m['qps_max']-m['qps'] for m in ms]],color=COL[name],label=NAMES[name],marker='o',ms=4,capsize=2,lw=1.8,alpha=1 if name in ['gt10_hard','linear','mlp','free'] else .7)
    if name=='gt10_hard':
     for xx,yy,m in zip(x,y,ms):ax.annotate(f"L{m['L']}",(xx,yy),xytext=(4,6),textcoords='offset points',fontsize=8,color=COL[name])
   ax.set(title=SHORT[ds],xlabel='Recall@10 (%)',ylabel='End-to-end QPS (higher is better)');ax.set_yscale('log');ax.set_ylim(300,4600);ax.yaxis.set_major_locator(matplotlib.ticker.FixedLocator([400,800,1000,2000,4000]));ax.yaxis.set_major_formatter(matplotlib.ticker.StrMethodFormatter('{x:,.0f}'))
  handles,labels=axs[0].get_legend_handles_labels();fig.legend(handles,labels,loc='lower center',bbox_to_anchor=(.5,.045),ncol=4 if len(names)==4 else 5,frameon=False)
  fig.suptitle('Offline shared entry transforms | fixed hubs, raw query, one HNSW top3',y=.98);fig.text(.5,.012,'Historical validation: 2,000 queries/dataset; 1 thread; median and range of 3 warm/plain walls. Lines connect tested L values.',ha='center',fontsize=9,color='#52606d');fig.subplots_adjust(left=.075,right=.985,bottom=.275,top=.86,wspace=.24);savefig(fig,fname)
 fig,axs=plt.subplots(1,2,figsize=(12,4.3))
 for ax,ds in zip(axs,DS):
  anchors=[100*m['recall_at_10'] for m in val[ds] if m['model_name']=='gt10_hard']
  for name in ['linear','mlp','free']:
   xs=sorted([r for r in interpolated if r['dataset']==SHORT[ds] and r['method']==name and any(abs(r['recall_target_pct']-a)<1e-8 for a in anchors)],key=lambda r:r['recall_target_pct'])
   x=[r['recall_target_pct'] for r in xs];y=[r['delta_qps_pct'] for r in xs];lo=[(r['ratio_min']-1)*100 for r in xs];hi=[(r['ratio_max']-1)*100 for r in xs]
   ax.plot(x,y,marker='o',color=COL[name],label=NAMES[name]);ax.fill_between(x,lo,hi,color=COL[name],alpha=.10)
  ax.axhline(0,color=COL['gt10_hard'],lw=.9);ax.set(title=SHORT[ds],xlabel='GT10 hard anchor Recall@10 (%)',ylabel='Interpolated QPS difference vs GT10 (%)')
 axs[1].legend(frameon=False);fig.suptitle('Same-Recall interpolation on tested curves | not extra measured points');fig.tight_layout();savefig(fig,'interpolated_speed')
 fig,axs=plt.subplots(1,2,figsize=(12,4.6))
 for ax,ds in zip(axs,DS):
  for j,name in enumerate(['linear','mlp','free']):
   cs=sorted([r for r in comps if r['dataset']==SHORT[ds] and r['method']==name],key=lambda r:r['L']);x=np.arange(5)+(j-1)*.12;y=np.array([r['delta_recall_pp'] for r in cs]);lo=np.array([r['ci_low_pp'] for r in cs]);hi=np.array([r['ci_high_pp'] for r in cs]);ax.errorbar(x,y,yerr=[y-lo,hi-y],label=NAMES[name],color=COL[name],marker='o',capsize=3,lw=1.3)
  ax.axhline(0,color='#243b53',lw=.8);ax.set_xticks(range(5),[64,128,256,400,800]);ax.set(title=SHORT[ds],xlabel='NSG search L',ylabel='Recall difference vs GT10 hard (pp)')
 axs[1].legend(frameon=False);fig.suptitle('Paired query uncertainty | 95% bootstrap intervals, unadjusted');fig.tight_layout();savefig(fig,'recall_differences')
 fig,axs=plt.subplots(2,2,figsize=(12,7))
 for col,ds in enumerate(DS):
  for family in ['linear','mlp','free']:
   for lam in [10,100]:
    curve=json.loads((O/'models'/ds/f'{family}_r{lam}'/'training_curve.json').read_text());style='-' if lam==10 else '--';label=f'{NAMES[family]}, reg={lam}'
    axs[0,col].plot([x['epoch'] for x in curve],[x['objective'] for x in curve],style,color=COL[family],label=label)
    axs[1,col].plot([x['epoch'] for x in curve],[x['drift'] for x in curve],style,color=COL[family])
  axs[0,col].set(title=SHORT[ds],ylabel='Training objective (lower is better)');axs[1,col].set(xlabel='Full-fit epoch',ylabel='Mean squared vector drift')
 axs[0,1].legend(fontsize=8,frameon=False);fig.suptitle('All 12 fits | same teacher, optimizer, epochs and regularization candidates');fig.tight_layout();savefig(fig,'training_curves')
if __name__=='__main__':main()
